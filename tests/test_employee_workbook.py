from decimal import Decimal
from dataclasses import replace
from io import BytesIO
from types import SimpleNamespace
import unittest

from openpyxl import Workbook

from payroll_tool.importers.employee_workbook import employee_from_row, parse_employee_workbook, yes
from payroll_tool.templates.hr_workbook import EXPECTED_HEADERS


class EmployeeWorkbookTests(unittest.TestCase):
    def test_damaged_xlsx_is_rejected_with_a_user_facing_error(self):
        with self.assertRaisesRegex(ValueError, "有效的 XLSX"):
            parse_employee_workbook(b"not-an-xlsx", {"full_time"})

    def test_employee_only_workbook_preserves_id_and_defaults_income_items(self):
        content = _workbook([
            {
                "employee_id": "0012",
                "employee_name": "测试员工",
                "employment_type": "intern",
                "intern_daily_rate": 200,
                "intern_hourly_rate": 25,
            }
        ])

        employees = parse_employee_workbook(content, {"intern"})

        self.assertEqual(employees[0].employee_id, "0012")
        self.assertEqual(employees[0].allowances, Decimal("0"))
        self.assertEqual(employees[0].bonuses, Decimal("0"))

    def test_numeric_id_uses_excel_zero_display_format(self):
        content = _workbook([
            {
                "employee_id": 12,
                "employee_id_number_format": "0000",
                "employee_name": "测试员工",
                "employment_type": "full_time",
                "monthly_salary": 10000,
            }
        ])

        employee = parse_employee_workbook(content, {"full_time"})[0]

        self.assertEqual(employee.employee_id, "0012")

    def test_duplicate_employee_ids_are_rejected(self):
        content = _workbook([
            {"employee_id": "0012", "employment_type": "intern"},
            {"employee_id": "0012", "employment_type": "intern"},
        ])

        with self.assertRaisesRegex(ValueError, "工号重复.*0012"):
            parse_employee_workbook(content, {"intern"})

    def test_employment_type_must_exist_in_published_policy(self):
        content = _workbook([
            {"employee_id": "0012", "employment_type": "contractor"}
        ])

        with self.assertRaisesRegex(
            ValueError, r"第 2 行.*工号 0012.*字段 employment_type.*contractor"
        ):
            parse_employee_workbook(content, {"full_time", "intern"})

    def test_policy_sheets_are_not_required(self):
        content = _workbook([
            {
                "employee_id": "A01",
                "employment_type": "full_time",
                "monthly_salary": 12000,
            }
        ])

        self.assertEqual(len(parse_employee_workbook(content, {"full_time"})), 1)

    def test_parser_reads_employee_cumulative_tax_fields(self):
        content = _workbook([{
            "employee_id": "S1", "employment_type": "full_time",
            "monthly_salary": 10000, "tax_year": 2026,
            "prior_taxable_income": 70000,
            "prior_special_deductions": 12000,
            "current_special_deductions": 1500,
            "prior_withheld_tax": 900,
        }])

        employee = parse_employee_workbook(content, {"full_time"})[0]

        self.assertEqual(employee.tax_year, 2026)
        self.assertEqual(employee.prior_taxable_income, Decimal("70000"))
        self.assertEqual(employee.prior_special_deductions, Decimal("12000"))
        self.assertEqual(employee.current_special_deductions, Decimal("1500"))
        self.assertEqual(employee.prior_withheld_tax, Decimal("900"))

    def test_invalid_tax_year_reports_row_employee_and_field(self):
        content = _workbook([{
            "employee_id": "S1", "employment_type": "full_time",
            "monthly_salary": 10000, "tax_year": "2026.5",
        }])

        with self.assertRaisesRegex(
            ValueError, r"第 2 行.*工号 S1.*字段 tax_year.*年份不是有效整数"
        ):
            parse_employee_workbook(content, {"full_time"})

    def test_invalid_hire_date_reports_row_employee_and_field(self):
        content = _workbook([{
            "employee_id": "S2", "employment_type": "full_time",
            "monthly_salary": 10000, "hire_date": "08/01/2026",
        }])

        with self.assertRaisesRegex(
            ValueError, r"第 2 行.*工号 S2.*字段 hire_date.*日期必须使用 YYYY-MM-DD"
        ):
            parse_employee_workbook(content, {"full_time"})

    def test_invalid_amount_reports_row_employee_field_and_reason(self):
        content = _workbook([{
            "employee_id": "S3", "employment_type": "full_time",
            "monthly_salary": "一万元",
        }])

        with self.assertRaisesRegex(
            ValueError, r"第 2 行.*工号 S3.*字段 monthly_salary.*金额不是有效数字"
        ):
            parse_employee_workbook(content, {"full_time"})

    def test_invalid_fixed_option_reports_field_context(self):
        content = _workbook([{
            "employee_id": "S4", "employment_type": "full_time",
            "monthly_salary": 10000, "statutory_enabled": "开启",
        }])

        with self.assertRaisesRegex(
            ValueError, r"第 2 行.*工号 S4.*字段 statutory_enabled.*是或否"
        ):
            parse_employee_workbook(content, {"full_time"})

    def test_shared_row_converter_keeps_legacy_option_semantics(self):
        employee = employee_from_row({
            "employee_id": "S5", "employment_type": "full_time",
            "monthly_salary": 10000, "statutory_enabled": "开启",
            "social_base_mode": "旧版自定义模式",
        })

        self.assertFalse(employee.statutory_enabled)
        self.assertEqual(employee.social_base_mode, "旧版自定义模式")
        self.assertFalse(yes("开启"))

    def test_boolean_employee_id_reports_row_and_field(self):
        content = _workbook([{
            "employee_id": True, "employment_type": "full_time",
            "monthly_salary": 10000,
        }])

        with self.assertRaisesRegex(
            ValueError, r"第 2 行.*字段 employee_id.*不能是布尔值"
        ):
            parse_employee_workbook(content, {"full_time"})

    def test_blank_statutory_switch_is_rejected(self):
        content = _workbook([{
            "employee_id": "S6", "employment_type": "full_time",
            "monthly_salary": 10000, "statutory_enabled": None,
        }])

        with self.assertRaisesRegex(
            ValueError, r"第 2 行.*工号 S6.*字段 statutory_enabled.*必填"
        ):
            parse_employee_workbook(content, {"full_time"})

    def test_negative_amount_is_rejected_with_field_context(self):
        content = _workbook([{
            "employee_id": "S7", "employment_type": "full_time",
            "monthly_salary": -1,
        }])

        with self.assertRaisesRegex(
            ValueError, r"第 2 行.*工号 S7.*字段 monthly_salary.*不能为负数"
        ):
            parse_employee_workbook(content, {"full_time"})

    def test_payroll_scheme_required_fields_use_published_policy(self):
        content = _workbook([{
            "employee_id": "S8", "employment_type": "intern",
            "intern_daily_rate": 200,
        }])
        policies = {
            "intern": SimpleNamespace(payroll_scheme="intern_daily_hourly")
        }

        with self.assertRaisesRegex(
            ValueError, r"第 2 行.*工号 S8.*字段 intern_hourly_rate.*必填"
        ):
            parse_employee_workbook(content, policies)

    def test_tier_code_must_exist_in_published_policy(self):
        content = _workbook([{
            "employee_id": "S9", "employment_type": "full_time",
            "monthly_salary": 10000, "statutory_enabled": "是",
            "pension_tier": "unknown",
        }])

        with self.assertRaisesRegex(
            ValueError, r"第 2 行.*工号 S9.*字段 pension_tier.*unknown"
        ):
            parse_employee_workbook(
                content,
                {"full_time"},
                lambda employee: {"pension": {"standard"}},
            )

    def test_legacy_set_call_does_not_validate_dynamic_tiers(self):
        content = _workbook([{
            "employee_id": "S10", "employment_type": "full_time",
            "monthly_salary": 10000, "pension_tier": "legacy-tier",
        }])

        employee = parse_employee_workbook(content, {"full_time"})[0]

        self.assertEqual(employee.pension_tier, "legacy-tier")

    def test_fixed_base_mode_requires_matching_amount(self):
        content = _workbook([{
            "employee_id": "S11", "employment_type": "full_time",
            "monthly_salary": 10000, "social_base_mode": "固定基数",
        }])

        with self.assertRaisesRegex(
            ValueError, r"第 2 行.*工号 S11.*字段 social_fixed_base.*必填"
        ):
            parse_employee_workbook(content, {"full_time"})

    def test_employee_transform_runs_before_tier_resolution(self):
        content = _workbook([{
            "employee_id": "S12", "employment_type": "full_time",
            "monthly_salary": 10000, "statutory_enabled": "是",
        }])
        resolved = []

        employees = parse_employee_workbook(
            content,
            {"full_time"},
            lambda employee: resolved.append(employee.statutory_enabled) or {},
            employee_transform=lambda employee: replace(
                employee, statutory_enabled=False
            ),
        )

        self.assertEqual(resolved, [False])
        self.assertFalse(employees[0].statutory_enabled)

    def test_employee_transform_tier_is_used_by_validation(self):
        content = _workbook([{
            "employee_id": "S13", "employment_type": "full_time",
            "monthly_salary": 10000, "statutory_enabled": "是",
        }])

        employee = parse_employee_workbook(
            content,
            {"full_time"},
            lambda employee: {"pension": {"standard", "special"}},
            employee_transform=lambda employee: replace(
                employee, pension_tier="standard"
            ),
        )[0]

        self.assertEqual(employee.pension_tier, "standard")

    def test_transformed_employment_type_controls_required_fields(self):
        content = _workbook([{
            "employee_id": "S14", "employment_type": "full_time",
            "monthly_salary": 10000,
        }])
        policies = {
            "full_time": SimpleNamespace(payroll_scheme="full_time_monthly"),
            "intern": SimpleNamespace(payroll_scheme="intern_daily_hourly"),
        }

        with self.assertRaisesRegex(
            ValueError, r"第 2 行.*工号 S14.*字段 intern_daily_rate.*必填"
        ):
            parse_employee_workbook(
                content,
                policies,
                employee_transform=lambda employee: replace(
                    employee, employment_type="intern"
                ),
            )


def _workbook(rows: list[dict[str, object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "员工薪酬"
    headers = EXPECTED_HEADERS["员工薪酬"]
    sheet.append(headers)
    for source in rows:
        values = {"statutory_enabled": "否", **source}
        sheet.append([values.get(header) for header in headers])
        if source.get("employee_id_number_format"):
            sheet.cell(sheet.max_row, 1).number_format = source["employee_id_number_format"]
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


if __name__ == "__main__":
    unittest.main()
