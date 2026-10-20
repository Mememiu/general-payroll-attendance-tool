from datetime import datetime, timezone
from io import BytesIO
from decimal import Decimal
import unittest

from openpyxl import load_workbook

from payroll_tool.audit import build_audit_manifest
from payroll_tool.exporters.excel import build_payroll_export


class ExportTests(unittest.TestCase):
    def test_export_contains_results_and_privacy_safe_audit_sheets(self):
        audit = build_audit_manifest(
            input_bytes=b"synthetic input",
            config_bytes=b"synthetic policy",
            result_bytes=b"synthetic result",
            timestamp=datetime(2026, 8, 14, tzinfo=timezone.utc),
            rule_package_id="policy",
            rule_version="1",
        )
        content = build_payroll_export(
            [{"employee_id": "SYN001", "base_pay": Decimal("10000.00")}],
            audit,
        )
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["计算结果", "审计信息"])
            self.assertEqual(workbook["计算结果"]["A2"].value, "SYN001")
            audit_values = [row[1].value for row in workbook["审计信息"].iter_rows(min_col=1, max_col=2)]
            self.assertIn("policy", audit_values)
            self.assertNotIn("SYN001", audit_values)
            self.assertNotIn(10000, audit_values)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
