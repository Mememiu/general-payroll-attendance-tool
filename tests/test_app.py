from io import BytesIO
from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from app import _apply_employee_exceptions, create_app
from desktop import LocalAppServer
from openpyxl import Workbook
from openpyxl import load_workbook
from payroll_tool.templates.hr_workbook import build_hr_workbook_template
from payroll_tool.templates.hr_workbook import EXPECTED_HEADERS
from payroll_tool.templates.employee_workbook import _validate_field_guidance
from tests.test_policy_validation import complete_policy
from payroll_tool.config.policy_serialization import EmployeeException
from payroll_tool.payroll.engine import EmployeePayProfile


class AppTests(unittest.TestCase):
    def test_housing_fund_exception_does_not_change_social_insurance_switch(self):
        employee = EmployeePayProfile(
            employee_id="S1", employment_type="full_time",
            monthly_salary=Decimal("10000"), statutory_enabled=True,
        )
        exception = EmployeeException(
            employee_id="S1", override_type="housing_fund_enabled",
            target_code="housing_fund", value=False,
            effective_from=date(2026, 8, 1), effective_to=None, reason="不缴公积金",
        )

        result = _apply_employee_exceptions(employee, (exception,))

        self.assertTrue(result.statutory_enabled)
        self.assertFalse(result.housing_fund_enabled)

    def setUp(self):
        self.app = create_app({"TESTING": True})
        self.client = self.app.test_client()

    def tearDown(self):
        self.app.extensions["task_store"].close()

    def test_desktop_server_binds_only_loopback(self):
        server = LocalAppServer(self.app, port=0)
        try:
            self.assertEqual(server.host, "127.0.0.1")
        finally:
            server.stop()

    def test_workspace_is_the_first_screen(self):
        response = self.client.get("/")

        self.assertEqual(response.status_code, 200)
        self.assertIn("本地薪酬绩效工具", response.get_data(as_text=True))
        self.assertIn("考勤文件", response.get_data(as_text=True))
        self.assertIn("员工薪酬文件", response.get_data(as_text=True))
        self.assertIn('name="policy_version_id"', response.get_data(as_text=True))
        self.assertIn('data-page="policies"', response.get_data(as_text=True))
        self.assertIn('type="month"', response.get_data(as_text=True))
        self.assertIn('name="employee_payroll"', response.get_data(as_text=True))
        self.assertNotIn('name="payroll_rules"', response.get_data(as_text=True))

    def test_calculation_records_are_read_only_and_do_not_expose_private_storage(self):
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        version = repository.publish(draft.id)
        for task_id in ("task-old", "task-new"):
            repository.create_calculation_record(
                task_id=task_id,
                policy_version_id=version.id,
                policy_digest=version.snapshot_sha256,
                payroll_month="2026-08",
                period_start="2026-08-01",
                period_end="2026-08-31",
                attendance_digest="a" * 64,
                employee_payroll_digest="b" * 64,
            )
        repository.complete_calculation_record(
            "task-new", employee_count=3, exception_count=1,
            result_available=True,
        )

        response = self.client.get("/api/calculation-records")

        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(
            [item["task_id"] for item in payload["items"]],
            ["task-new", "task-old"],
        )
        self.assertTrue(payload["items"][0]["result_available"])
        serialized = response.get_data(as_text=True).lower()
        for private_name in ("token", "path", "download"):
            self.assertNotIn(private_name, serialized)

    def test_task_with_published_policy_writes_completed_calculation_record(self):
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        draft = repository.save_draft(draft.id, complete_policy())
        version = repository.publish(draft.id)

        response = self.client.post(
            "/tasks",
            data={
                "policy_version_id": version.id,
                "payroll_month": "2026-08",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
                "attendance": (_attendance_workbook(), "attendance.xlsx"),
                "employee_payroll": (_employee_only_workbook(), "employees.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201)
        record = repository.get_calculation_record(response.get_json()["task_id"])
        self.assertEqual(record.status, "completed")
        self.assertEqual(record.policy_version_id, version.id)
        self.assertEqual(record.policy_digest, version.snapshot_sha256)
        self.assertEqual(record.employee_count, 1)
        self.assertTrue(record.result_available)
        self.assertEqual(len(record.attendance_digest), 64)
        self.assertEqual(len(record.employee_payroll_digest), 64)

    def test_page_two_file_request_uses_published_policy_snapshot(self):
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        draft = repository.save_draft(draft.id, complete_policy())
        version = repository.publish(draft.id)

        response = self.client.post(
            "/tasks",
            data={
                "policy_version_id": version.id,
                "payroll_month": "2026-08",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
                "attendance": (_attendance_workbook(), "attendance.xlsx"),
                "employee_payroll": (_employee_only_workbook(), "employees.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201, response.get_json())
        record = repository.get_calculation_record(response.get_json()["task_id"])
        self.assertEqual(record.policy_digest, version.snapshot_sha256)

    def test_published_policy_calculates_tax_without_prior_tax_workbook(self):
        policy = complete_policy()
        policy["statutory"]["tax"].update(
            monthly_exemption="0",
            brackets=[{"threshold": "0", "rate": "0.10", "quick_deduction": "0"}],
        )
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        draft = repository.save_draft(draft.id, policy)
        version = repository.publish(draft.id)

        response = self.client.post(
            "/tasks",
            data={
                "policy_version_id": version.id,
                "payroll_month": "2026-08",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
                "attendance": (_attendance_workbook(), "attendance.xlsx"),
                "employee_payroll": (_employee_only_workbook(), "employees.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201, response.get_json())
        payload = response.get_json()
        download = self.client.get(
            f"/tasks/{payload['task_id']}/download?token={payload['download_token']}"
        )
        workbook = load_workbook(BytesIO(download.data), read_only=True, data_only=True)
        try:
            headers = [cell.value for cell in workbook["计算结果"][1]]
            values = [cell.value for cell in workbook["计算结果"][2]]
            self.assertGreater(dict(zip(headers, values))["individual_income_tax"], 0)
        finally:
            workbook.close()

    def test_published_policy_uses_employee_cumulative_tax_fields(self):
        policy = complete_policy()
        policy["statutory"]["tax"].update(
            brackets=[{"threshold": "0", "rate": "0.10", "quick_deduction": "0"}],
        )
        employee_file = _employee_only_workbook()
        workbook = load_workbook(employee_file)
        headers = [cell.value for cell in workbook["员工薪酬"][1]]
        values = {
            "tax_year": 2026, "prior_taxable_income": 100000,
            "prior_special_deductions": 0, "prior_withheld_tax": 1000,
        }
        for name, value in values.items():
            workbook["员工薪酬"].cell(2, headers.index(name) + 1).value = value
        updated = BytesIO()
        workbook.save(updated)
        workbook.close()
        updated.seek(0)

        result = self._run_published_policy_and_read_result(policy, updated)

        self.assertGreater(result["individual_income_tax"], 0)

    def test_employment_type_can_disable_tax_for_published_policy(self):
        policy = complete_policy()
        policy["statutory"]["tax"].update(
            monthly_exemption="0",
            brackets=[{"threshold": "0", "rate": "0.10", "quick_deduction": "0"}],
        )
        policy["employment_types"][0]["tax_applicable"] = False

        result = self._run_published_policy_and_read_result(policy)

        self.assertEqual(result["individual_income_tax"], 0)

    def test_tax_is_not_applied_before_hr_effective_date(self):
        policy = complete_policy()
        policy["statutory"]["tax"].update(
            effective_from="2027-01-01",
            monthly_exemption="0",
            brackets=[{"threshold": "0", "rate": "0.10", "quick_deduction": "0"}],
        )

        result = self._run_published_policy_and_read_result(policy)

        self.assertEqual(result["individual_income_tax"], 0)

    def test_tax_effective_on_manual_period_start_is_applied(self):
        policy = complete_policy()
        policy["statutory"]["tax"].update(
            effective_from="2026-08-15", monthly_exemption="0",
            brackets=[{"threshold": "0", "rate": "0.10", "quick_deduction": "0"}],
        )
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-15")
        draft = repository.save_draft(draft.id, policy)
        version = repository.publish(draft.id)
        response = self.client.post(
            "/tasks",
            data={
                "policy_version_id": version.id, "payroll_month": "2026-08",
                "period_start": "2026-08-15", "period_end": "2026-08-31",
                "attendance": (_attendance_workbook_for("2026-08-15"), "attendance.xlsx"),
                "employee_payroll": (_employee_only_workbook(), "employees.xlsx"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 201, response.get_json())
        payload = response.get_json()
        download = self.client.get(
            f"/tasks/{payload['task_id']}/download?token={payload['download_token']}"
        )
        workbook = load_workbook(BytesIO(download.data), read_only=True, data_only=True)
        try:
            headers = [cell.value for cell in workbook["计算结果"][1]]
            values = [cell.value for cell in workbook["计算结果"][2]]
            self.assertGreater(dict(zip(headers, values))["individual_income_tax"], 0)
        finally:
            workbook.close()

    def test_employment_type_can_disable_all_statutory_contributions(self):
        policy = complete_policy()
        policy["employment_types"][0].update(
            social_insurance_applicable=False,
            housing_fund_applicable=False,
        )
        policy["statutory"]["insurance"] = [{
            "code": "pension", "name": "养老保险", "tier_code": "standard",
            "enabled": True, "region": "测试地区", "effective_from": "2026-01-01",
            "effective_to": None, "policy_source": "HR 手工录入",
            "employment_types": ["full_time"], "hire_month_rule": "always",
            "termination_month_rule": "always", "base_mode": "fixed",
            "base_lower": "5000", "base_upper": "5000", "fixed_base": "5000",
            "employee_rate": "0.08", "employer_rate": "0.16",
        }]
        employee_file = _employee_only_workbook()
        workbook = load_workbook(employee_file)
        headers = [cell.value for cell in workbook["员工薪酬"][1]]
        workbook["员工薪酬"].cell(2, headers.index("statutory_enabled") + 1).value = "是"
        updated = BytesIO()
        workbook.save(updated)
        workbook.close()
        updated.seek(0)

        result = self._run_published_policy_and_read_result(policy, updated)

        self.assertEqual(result["employee_statutory_deduction"], 0)
        self.assertEqual(result["employer_statutory_cost"], 0)

    def _run_published_policy_and_read_result(
        self, policy, employee_file=None, attendance_file=None
    ):
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        draft = repository.save_draft(draft.id, policy)
        version = repository.publish(draft.id)
        response = self.client.post(
            "/tasks",
            data={
                "policy_version_id": version.id,
                "payroll_month": "2026-08",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
                "attendance": (
                    attendance_file or _attendance_workbook(), "attendance.xlsx"
                ),
                "employee_payroll": (employee_file or _employee_only_workbook(), "employees.xlsx"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 201, response.get_json())
        payload = response.get_json()
        download = self.client.get(
            f"/tasks/{payload['task_id']}/download?token={payload['download_token']}"
        )
        workbook = load_workbook(BytesIO(download.data), read_only=True, data_only=True)
        try:
            headers = [cell.value for cell in workbook["计算结果"][1]]
            values = [cell.value for cell in workbook["计算结果"][2]]
            return dict(zip(headers, values))
        finally:
            workbook.close()

    def test_configured_overtime_is_settled_and_exported(self):
        policy = _cash_and_comp_overtime_policy()
        expected_pay = {
            "工作日": 192.03,
            "休息日": 256.04,
            "法定节假日": 384.06,
        }

        for workday_type, overtime_pay in expected_pay.items():
            with self.subTest(workday_type=workday_type):
                result = self._run_published_policy_and_read_result(
                    policy,
                    attendance_file=_overtime_attendance_workbook(workday_type),
                )
                self.assertEqual(result["raw_overtime_hours"], 2)
                self.assertEqual(result["overtime_hours"], 2)
                self.assertEqual(result["comp_time_accrual_minutes"], 180)
                self.assertEqual(result["overtime_pay"], overtime_pay)
                self.assertEqual(
                    round(result["gross_pay"] - result["base_pay"], 2),
                    overtime_pay,
                )
                self.assertIsNone(result["review_reasons"])

    def test_cash_overtime_without_valid_workday_type_needs_review(self):
        policy = _cash_and_comp_overtime_policy()
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        draft = repository.save_draft(draft.id, policy)
        version = repository.publish(draft.id)

        response = self.client.post(
            "/tasks",
            data={
                "policy_version_id": version.id,
                "payroll_month": "2026-08",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
                "attendance": (
                    _overtime_attendance_workbook("workday"), "attendance.xlsx"
                ),
                "employee_payroll": (_employee_only_workbook(), "employees.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201, response.get_json())
        payload = response.get_json()
        self.assertGreater(payload["review_item_count"], 0)
        self.assertEqual(
            repository.get_calculation_record(payload["task_id"]).status,
            "needs_review",
        )
        download = self.client.get(
            f"/tasks/{payload['task_id']}/download?token={payload['download_token']}"
        )
        workbook = load_workbook(BytesIO(download.data), read_only=True, data_only=True)
        try:
            headers = [cell.value for cell in workbook["计算结果"][1]]
            values = [cell.value for cell in workbook["计算结果"][2]]
            result = dict(zip(headers, values))
            self.assertEqual(result["overtime_hours"], 2)
            self.assertEqual(result["overtime_pay"], 0)
            self.assertIn("工作日类型", result["review_reasons"])
        finally:
            workbook.close()

    def test_cash_disabled_does_not_require_workday_type_column(self):
        policy = _cash_and_comp_overtime_policy()
        policy["company"]["overtime"]["cash_enabled"] = False
        for key in ("cash_formula", "cash_rate", "day_rates"):
            policy["company"]["overtime"].pop(key)

        result = self._run_published_policy_and_read_result(
            policy,
            attendance_file=_overtime_attendance_workbook(),
        )

        self.assertEqual(result["overtime_hours"], 2)
        self.assertEqual(result["overtime_pay"], 0)

    def test_overtime_overlapping_leave_exports_review_without_settlement(self):
        policy = _cash_and_comp_overtime_policy()
        policy["company"]["leaves"].append({
            "code": "personal_leave", "name": "事假", "category": "leave",
            "pay_scheme": "unpaid", "counts_as_attendance": False,
            "counts_as_absence": True, "affects_perfect_attendance": True,
            "skip_punch_requirement": False, "proof_required": False,
            "missing_proof_action": "review", "cross_day_allowed": False,
            "deduction_category": "attendance", "enabled": True,
        })

        result = self._run_published_policy_and_read_result(
            policy,
            attendance_file=_overtime_attendance_workbook(
                "工作日",
                attendance_text=(
                    "正常(09:30),正常(20:00);加班(18:00-20:00);"
                    "事假(19:00-19:30)"
                ),
            ),
        )

        self.assertEqual(result["raw_overtime_hours"], 2)
        self.assertEqual(result["overtime_hours"], 0)
        self.assertEqual(result["comp_time_accrual_minutes"], 0)
        self.assertEqual(result["overtime_pay"], 0)
        self.assertIn("与非加班审批重叠", result["review_reasons"])

    def test_two_uploaded_files_with_same_name_do_not_overwrite_each_other(self):
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        draft = repository.save_draft(draft.id, complete_policy())
        version = repository.publish(draft.id)

        response = self.client.post(
            "/tasks",
            data={
                "policy_version_id": version.id,
                "payroll_month": "2026-08",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
                "attendance": (_attendance_workbook(), "8月数据.xlsx"),
                "employee_payroll": (_employee_only_workbook(), "8月数据.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201, response.get_json())

    def test_result_write_failure_transitions_record_to_failed(self):
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        draft = repository.save_draft(draft.id, complete_policy())
        version = repository.publish(draft.id)
        with patch.object(
            self.app.extensions["task_store"], "save_result",
            side_effect=OSError("disk full"),
        ):
            with self.assertRaises(OSError):
                self.client.post(
                    "/tasks",
                    data={
                        "policy_version_id": version.id,
                        "payroll_month": "2026-08",
                        "period_start": "2026-08-01",
                        "period_end": "2026-08-31",
                        "attendance": (_attendance_workbook(), "attendance.xlsx"),
                        "employee_payroll": (_employee_only_workbook(), "employees.xlsx"),
                    },
                    content_type="multipart/form-data",
                )
        record = repository.list_calculation_records()[0]
        self.assertEqual(record.status, "failed")
        self.assertFalse(record.result_available)

    def test_completed_result_remains_downloadable_after_app_restart(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            config = {
                "TESTING": True,
                "POLICY_DATABASE": root / "policies.sqlite3",
                "TASK_ROOT": root / "tasks",
            }
            first = create_app(config)
            repository = first.extensions["policy_repository"]
            draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
            draft = repository.save_draft(draft.id, complete_policy())
            version = repository.publish(draft.id)
            response = first.test_client().post(
                "/tasks",
                data={
                    "policy_version_id": version.id,
                    "payroll_month": "2026-08",
                    "period_start": "2026-08-01",
                    "period_end": "2026-08-31",
                    "attendance": (_attendance_workbook(), "attendance.xlsx"),
                    "employee_payroll": (_employee_only_workbook(), "employees.xlsx"),
                },
                content_type="multipart/form-data",
            )
            self.assertEqual(response.status_code, 201, response.get_json())
            payload = response.get_json()
            first.extensions["task_store"].close()

            restarted = create_app(config)
            client = restarted.test_client()
            denied = client.get(
                f"/tasks/{payload['task_id']}/download?token=wrong"
            )
            download = client.get(
                f"/tasks/{payload['task_id']}/download?token={payload['download_token']}"
            )

            self.assertEqual(denied.status_code, 403)
            self.assertEqual(download.status_code, 200)
            self.assertGreater(len(download.data), 1000)
            restarted.extensions["task_store"].close()

    def test_task_with_attendance_review_items_is_marked_needs_review(self):
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        draft = repository.save_draft(draft.id, complete_policy())
        version = repository.publish(draft.id)
        attendance = _attendance_workbook()
        workbook = load_workbook(attendance)
        workbook.active["C2"] = "正常(09:00),缺卡(-)"
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        content.seek(0)

        response = self.client.post(
            "/tasks",
            data={
                "policy_version_id": version.id,
                "payroll_month": "2026-08",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
                "attendance": (content, "attendance.xlsx"),
                "employee_payroll": (_employee_only_workbook(), "employees.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201, response.get_json())
        record = repository.get_calculation_record(response.get_json()["task_id"])
        self.assertEqual(record.status, "needs_review")
        self.assertGreater(record.exception_count, 0)

    def test_active_employee_exception_is_applied_and_counted(self):
        policy = complete_policy()
        policy["company"]["payroll"]["perfect_attendance_award"] = "200"
        policy["employee_exceptions"] = [{
            "employee_id": "SYN001",
            "override_type": "perfect_attendance_amount",
            "target_code": "perfect_attendance",
            "value": "0",
            "effective_from": "2026-08-01",
            "effective_to": "2026-08-31",
            "reason": "当月取消",
            "enabled": True,
        }]
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        draft = repository.save_draft(draft.id, policy)
        version = repository.publish(draft.id)

        response = self.client.post(
            "/tasks",
            data={
                "policy_version_id": version.id,
                "payroll_month": "2026-08",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
                "attendance": (_attendance_workbook(), "attendance.xlsx"),
                "employee_payroll": (_employee_only_workbook(), "employees.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201, response.get_json())
        payload = response.get_json()
        record = repository.get_calculation_record(payload["task_id"])
        self.assertEqual(record.exception_count, 1)
        self.assertEqual(record.status, "completed")
        download = self.client.get(
            f"/tasks/{payload['task_id']}/download?token={payload['download_token']}"
        )
        workbook = load_workbook(BytesIO(download.data), read_only=True, data_only=True)
        try:
            headers = [cell.value for cell in workbook["计算结果"][1]]
            values = [cell.value for cell in workbook["计算结果"][2]]
            self.assertEqual(dict(zip(headers, values))["perfect_attendance_award"], 0)
        finally:
            workbook.close()

    def test_mid_period_employment_type_exception_requires_split_calculation(self):
        policy = complete_policy()
        policy["employee_exceptions"] = [{
            "employee_id": "SYN001", "override_type": "employment_type",
            "target_code": "employment_type", "value": "full_time",
            "effective_from": "2026-08-15", "effective_to": None,
            "reason": "月中变更", "enabled": True,
        }]
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        draft = repository.save_draft(draft.id, policy)
        version = repository.publish(draft.id)

        response = self.client.post(
            "/tasks",
            data={
                "policy_version_id": version.id,
                "payroll_month": "2026-08",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
                "attendance": (_attendance_workbook(), "attendance.xlsx"),
                "employee_payroll": (_employee_only_workbook(), "employees.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 400)
        self.assertTrue(any(
            "用工类型例外" in error and "完整计薪周期" in error
            for error in response.get_json()["errors"]
        ))

    def test_commercial_insurance_employee_exception_reaches_export(self):
        policy = complete_policy()
        policy["commercial_insurance"] = [{
            "code": "accident", "name": "团体意外险", "enabled": True,
            "region": "测试地区", "effective_from": "2026-01-01",
            "effective_to": None, "calculation": "fixed_amount", "amount": "100",
            "employee_share": "0", "employer_share": "1",
            "taxable": False, "included_in_gross": False,
            "deduction_mode": "company_paid", "employment_types": ["full_time"],
        }]
        policy["employee_exceptions"] = [{
            "employee_id": "SYN001", "override_type": "commercial_amount",
            "target_code": "accident", "value": "150",
            "effective_from": "2026-08-01", "effective_to": "2026-08-31",
            "reason": "员工保额不同", "enabled": True,
        }]

        result = self._run_published_policy_and_read_result(policy)

        self.assertEqual(result["employee_commercial_insurance_deduction"], 0)
        self.assertEqual(result["employer_commercial_insurance_cost"], 150)

    def test_task_validation_failure_keeps_failed_calculation_record(self):
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        draft = repository.save_draft(draft.id, complete_policy())
        version = repository.publish(draft.id)

        response = self.client.post(
            "/tasks",
            data={
                "policy_version_id": version.id,
                "payroll_month": "2026-08",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
                "attendance": (BytesIO(b"not-an-xlsx"), "attendance.xlsx"),
                "employee_payroll": (_employee_only_workbook(), "employees.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 400)
        records = repository.list_calculation_records()
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0].status, "failed")
        self.assertFalse(records[0].result_available)
        self.assertTrue(records[0].error_summary)

    def test_unexpected_task_error_transitions_record_to_failed(self):
        repository = self.app.extensions["policy_repository"]
        draft = repository.create_draft("公司制度", "2026.08", "2026-08-01")
        draft = repository.save_draft(draft.id, complete_policy())
        version = repository.publish(draft.id)

        with patch("app.import_attendance_workbook", side_effect=RuntimeError("boom")):
            with self.assertRaises(RuntimeError):
                self.client.post(
                    "/tasks",
                    data={
                        "policy_version_id": version.id,
                        "payroll_month": "2026-08",
                        "period_start": "2026-08-01",
                        "period_end": "2026-08-31",
                        "attendance": (_attendance_workbook(), "attendance.xlsx"),
                        "employee_payroll": (_employee_only_workbook(), "employees.xlsx"),
                    },
                    content_type="multipart/form-data",
                )

        record = repository.list_calculation_records()[0]
        self.assertEqual(record.status, "failed")
        self.assertEqual(record.error_summary, "内部计算错误")

    def test_task_requires_attendance_payroll_and_policy_uploads(self):
        response = self.client.post(
            "/tasks",
            data={"attendance": (BytesIO(b"attendance"), "attendance.xlsx")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 400)
        self.assertTrue(any("薪酬规则工作簿" in error for error in response.get_json()["errors"]))

    def test_task_rejects_reversed_payroll_period(self):
        response = self.client.post(
            "/tasks",
            data={
                "payroll_month": "2026-08",
                "period_start": "2026-08-31",
                "period_end": "2026-08-01",
                "attendance": (_attendance_workbook(), "attendance.xlsx"),
                "payroll_rules": (_completed_hr_workbook(), "payroll-rules.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 400)
        self.assertTrue(any("起薪日" in error for error in response.get_json()["errors"]))

    def test_task_uses_current_request_uploads_and_returns_private_task_id(self):
        response = self.client.post(
            "/tasks",
            data={
                "payroll_month": "2026-08",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
                "attendance": (_attendance_workbook(), "attendance.xlsx"),
                "payroll_rules": (_completed_hr_workbook(), "payroll-rules.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201)
        payload = response.get_json()
        self.assertEqual(len(payload["task_id"]), 32)
        self.assertNotIn("attendance", payload["task_id"])
        self.assertEqual(payload["employee_count"], 1)
        self.assertEqual(payload["attendance_row_count"], 1)
        self.assertEqual(payload["processed_day_count"], 1)
        self.assertEqual(payload["review_item_count"], 0)
        self.assertEqual(payload["payroll_result_count"], 1)

        denied = self.client.get(f"/tasks/{payload['task_id']}/download?token=wrong")
        self.assertEqual(denied.status_code, 403)

        download = self.client.get(
            f"/tasks/{payload['task_id']}/download?token={payload['download_token']}"
        )
        self.assertEqual(download.status_code, 200)
        workbook = load_workbook(BytesIO(download.data), read_only=True, data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["计算结果", "审计信息"])
            self.assertEqual(workbook["计算结果"]["A2"].value, "SYN001")
            self.assertIn("employee_name", [cell.value for cell in workbook["计算结果"][1]])
            headers = [cell.value for cell in workbook["计算结果"][1]]
            self.assertIn("attendance_days", headers)
            self.assertIn("effective_hours", headers)
            self.assertIn("absence_hours", headers)
            self.assertIn("late_hours", headers)
            self.assertIn("leave_hours", headers)
            self.assertIn("overtime_hours", headers)
            self.assertIn("overtime_pay", headers)
        finally:
            workbook.close()

    def test_task_rejects_blank_hr_workbook_with_missing_required_values(self):
        response = self.client.post(
            "/tasks",
            data={
                "payroll_month": "2026-08",
                "period_start": "2026-08-01",
                "period_end": "2026-08-31",
                "attendance": (BytesIO(b"attendance"), "attendance.xlsx"),
                "payroll_rules": (BytesIO(build_hr_workbook_template()), "payroll-rules.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 400)
        self.assertTrue(any("必填" in error for error in response.get_json()["errors"]))

    def test_task_applies_employee_cumulative_tax_to_downloaded_result(self):
        rules = _completed_hr_workbook()
        workbook = load_workbook(rules)
        employee_sheet = workbook["员工薪酬"]
        employee_row = next(
            row for row in range(2, employee_sheet.max_row + 1)
            if employee_sheet.cell(row, 1).value == "SYN001"
        )
        employee_sheet.cell(employee_row, 8).value = 21750
        workbook["个税累计"].append(["SYN001", 2026, 0, 0, 0, ""])
        updated_rules = BytesIO()
        workbook.save(updated_rules)
        workbook.close()
        updated_rules.seek(0)

        response = self.client.post(
            "/tasks",
            data={
                "payroll_month": "2026-01",
                "period_start": "2026-01-01",
                "period_end": "2026-01-31",
                "attendance": (_attendance_workbook_for("2026-01-03", days=21), "attendance.xlsx"),
                "payroll_rules": (updated_rules, "payroll-rules.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201)
        payload = response.get_json()
        download = self.client.get(
            f"/tasks/{payload['task_id']}/download?token={payload['download_token']}"
        )
        workbook = load_workbook(BytesIO(download.data), read_only=True, data_only=True)
        try:
            headers = [cell.value for cell in workbook["计算结果"][1]]
            values = [cell.value for cell in workbook["计算结果"][2]]
            result = dict(zip(headers, values, strict=True))
            self.assertEqual(result["individual_income_tax"], 480)
            self.assertEqual(result["net_pay"], 20520)
        finally:
            workbook.close()

    def test_task_calculates_current_tax_when_employee_has_no_prior_record(self):
        rules = _completed_hr_workbook()
        workbook = load_workbook(rules)
        employee_sheet = workbook["员工薪酬"]
        employee_row = next(
            row for row in range(2, employee_sheet.max_row + 1)
            if employee_sheet.cell(row, 1).value == "SYN001"
        )
        employee_sheet.cell(employee_row, 8).value = 21750
        updated_rules = BytesIO()
        workbook.save(updated_rules)
        workbook.close()
        updated_rules.seek(0)

        response = self.client.post(
            "/tasks",
            data={
                "payroll_month": "2026-01",
                "period_start": "2026-01-01",
                "period_end": "2026-01-31",
                "attendance": (_attendance_workbook_for("2026-01-03", days=21), "attendance.xlsx"),
                "payroll_rules": (updated_rules, "payroll-rules.xlsx"),
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201)
        payload = response.get_json()
        download = self.client.get(
            f"/tasks/{payload['task_id']}/download?token={payload['download_token']}"
        )
        workbook = load_workbook(BytesIO(download.data), read_only=True, data_only=True)
        try:
            headers = [cell.value for cell in workbook["计算结果"][1]]
            values = [cell.value for cell in workbook["计算结果"][2]]
            result = dict(zip(headers, values, strict=True))
            self.assertEqual(result["individual_income_tax"], 480)
        finally:
            workbook.close()

    def test_hr_workbook_template_download_contains_fixed_sheets(self):
        response = self.client.get("/templates/payroll-rules.xlsx")

        self.assertEqual(response.status_code, 200)
        self.assertGreater(len(response.data), 5000)

    def test_employee_payroll_template_is_blank_and_matches_import_columns(self):
        response = self.client.get("/templates/employee-payroll.xlsx")

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data), data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["填写说明", "字段说明", "员工薪酬"])
            guidance = workbook["字段说明"]
            guidance_rows = list(guidance.iter_rows(min_row=2, values_only=True))
            self.assertEqual(
                [row[0] for row in guidance_rows], EXPECTED_HEADERS["员工薪酬"]
            )
            self.assertTrue(all(len(row) == 6 and all(row) for row in guidance_rows))
            requirements = {row[0]: row[2] for row in guidance_rows}
            for field in (
                "prior_taxable_income",
                "prior_special_deductions",
                "prior_withheld_tax",
            ):
                self.assertEqual(requirements[field], "条件必填")
            guide_text = "\n".join(
                str(cell.value or "") for row in workbook["填写说明"] for cell in row
            )
            self.assertIn("填写流程", guide_text)
            self.assertIn("本机", guide_text)
            headers = [cell.value for cell in workbook["员工薪酬"][1]]
            self.assertEqual(headers, EXPECTED_HEADERS["员工薪酬"])
            for field in ("employee_id", "statutory_enabled", "employment_type"):
                cell = workbook["员工薪酬"].cell(1, headers.index(field) + 1)
                self.assertIsNotNone(cell.comment)
                self.assertEqual(cell.comment.text.splitlines()[1], "必填")
                self.assertEqual(cell.fill.fgColor.rgb, "00F8D7DA")
            conditional = workbook["员工薪酬"].cell(
                1, headers.index("monthly_salary") + 1
            )
            optional = workbook["员工薪酬"].cell(1, headers.index("employee_name") + 1)
            self.assertEqual(conditional.fill.fgColor.rgb, "00FFF3CD")
            self.assertEqual(optional.fill.fgColor.rgb, "00DDEBF7")
            self.assertTrue(
                all(
                    cell.value is None
                    for row in workbook["员工薪酬"].iter_rows(min_row=2)
                    for cell in row
                )
            )
            self.assertEqual(workbook["员工薪酬"]["A2"].number_format, "@")
            self.assertGreaterEqual(len(workbook["员工薪酬"].data_validations.dataValidation), 3)
            statutory_validation = next(
                validation
                for validation in workbook["员工薪酬"].data_validations.dataValidation
                if "B2:B1001" in str(validation.sqref)
            )
            self.assertFalse(statutory_validation.allow_blank)
        finally:
            workbook.close()

    def test_employee_field_guidance_rejects_schema_drift(self):
        with self.assertRaisesRegex(ValueError, "字段说明.*标准表头"):
            _validate_field_guidance(("employee_id",), (("different",),))

    def test_rule_pack_can_be_saved_listed_and_downloaded(self):
        with TemporaryDirectory() as directory:
            app = create_app({"TESTING": True, "RULE_PACK_ROOT": directory})
            client = app.test_client()
            workbook = _completed_hr_workbook().getvalue()
            response = client.post(
                "/rule-packs",
                data={
                    "name": "Synthetic policy",
                    "version": "2026.08",
                    "effective_from": "2026-08-01",
                    "workbook": (BytesIO(workbook), "rules.xlsx"),
                },
                content_type="multipart/form-data",
            )

            self.assertEqual(response.status_code, 201)
            pack_id = response.get_json()["id"]
            self.assertEqual(client.get("/rule-packs").get_json()["items"][0]["id"], pack_id)
            self.assertEqual(client.get(f"/rule-packs/{pack_id}/download").data, workbook)
            app.extensions["task_store"].close()

    def test_task_can_use_saved_rule_pack_without_reuploading_workbook(self):
        with TemporaryDirectory() as directory:
            app = create_app({"TESTING": True, "RULE_PACK_ROOT": directory})
            client = app.test_client()
            rules = _completed_hr_workbook().getvalue()
            saved = client.post(
                "/rule-packs",
                data={
                    "name": "Synthetic policy", "version": "2026.08",
                    "effective_from": "2026-08-01",
                    "workbook": (BytesIO(rules), "rules.xlsx"),
                }, content_type="multipart/form-data",
            ).get_json()
            response = client.post(
                "/tasks",
                data={
                    "payroll_month": "2026-08", "period_start": "2026-08-01",
                    "period_end": "2026-08-31", "rule_pack_id": saved["id"],
                    "attendance": (_attendance_workbook(), "attendance.xlsx"),
                }, content_type="multipart/form-data",
            )

            self.assertEqual(response.status_code, 201)
            app.extensions["task_store"].close()

    def test_rule_pack_clone_requires_new_valid_workbook(self):
        with TemporaryDirectory() as directory:
            app = create_app({"TESTING": True, "RULE_PACK_ROOT": directory})
            client = app.test_client()
            workbook = _completed_hr_workbook().getvalue()
            original = client.post(
                "/rule-packs",
                data={"name": "Synthetic policy", "version": "2026.01", "effective_from": "2026-01-01", "workbook": (BytesIO(workbook), "rules.xlsx")},
                content_type="multipart/form-data",
            ).get_json()
            clone = client.post(
                f"/rule-packs/{original['id']}/clone",
                data={"version": "2026.08", "effective_from": "2026-08-01", "workbook": (BytesIO(workbook), "updated-rules.xlsx")},
                content_type="multipart/form-data",
            )

            self.assertEqual(clone.status_code, 201)
            self.assertEqual(clone.get_json()["version"], "2026.08")
            app.extensions["task_store"].close()


def _completed_hr_workbook() -> BytesIO:
    from openpyxl import load_workbook

    content = BytesIO(build_hr_workbook_template())
    workbook = load_workbook(content)
    base = workbook["基础与考勤"]
    values = {
        "package_id": "synthetic-policy",
        "policy_version": "1",
        "work_mode": "固定工时",
        "overtime_start_mode": "审批开始时间",
        "daily_required_hours": 8,
        "monthly_pay_days": 21.75,
        "standard_start": "09:00",
        "standard_end": "18:30",
        "lunch_start": "12:00",
        "lunch_end": "13:30",
        "earliest_valid_start": "08:00",
        "latest_normal_start": "09:15",
        "late_grace_end": "09:30",
        "late_makeup_allowed": "否",
        "full_day_absence_after_hours": 2,
        "attendance_deduction_factor": 1,
        "sick_leave_deduction_factor": 0.3,
        "perfect_attendance_award": 0,
        "overtime_cash_enabled": "否",
        "entry_day_full_attendance": "是",
        "termination_day_full_attendance": "是",
        "allowance_taxable": "是",
        "allowance_included_in_social_base": "否",
        "allowance_included_in_fund_base": "否",
        "bonus_taxable": "是",
        "bonus_included_in_social_base": "否",
        "bonus_included_in_fund_base": "否",
    }
    for row in range(2, base.max_row + 1):
        code = base.cell(row, 1).value
        if code in values:
            base.cell(row, 3).value = values[code]
    mapping = workbook["考勤字段映射"]
    mapping["D2"], mapping["D3"], mapping["D4"] = "工号", "日期", "每日考勤结果"
    leave = workbook["假勤与豁免"]
    for row in range(2, leave.max_row + 1):
        for column, value in {3: 1, 4: "否", 5: "是", 6: "是", 7: "是", 8: "否"}.items():
            leave.cell(row, column).value = value
    social = workbook["五险一金"]
    for row in range(2, social.max_row + 1):
        social.cell(row, 4).value = "否"
    employee = workbook["员工薪酬"]
    employee.append(["SYN001", "否", "Synthetic User", "正职", None, None, None, 10000])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def _attendance_workbook() -> BytesIO:
    return _attendance_workbook_for("2026-08-01")


def _employee_only_workbook() -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "员工薪酬"
    sheet.append(EXPECTED_HEADERS["员工薪酬"])
    values = {
        "employee_id": "SYN001", "employee_name": "Synthetic User",
        "employment_type": "full_time", "monthly_salary": 10000,
        "statutory_enabled": "否",
    }
    sheet.append([values.get(header) for header in EXPECTED_HEADERS["员工薪酬"]])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def _attendance_workbook_for(work_date: str, *, days: int = 1) -> BytesIO:
    from datetime import date, timedelta

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["工号", "日期", "每日考勤结果"])
    start = date.fromisoformat(work_date)
    for offset in range(days):
        sheet.append([
            "SYN001",
            (start + timedelta(days=offset)).isoformat(),
            "正常(09:00),正常(18:30)",
        ])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def _cash_and_comp_overtime_policy() -> dict[str, object]:
    policy = complete_policy()
    policy["company"]["leaves"].append({
        "code": "overtime",
        "name": "加班",
        "category": "overtime",
        "pay_scheme": "full_pay",
        "counts_as_attendance": True,
        "counts_as_absence": False,
        "affects_perfect_attendance": False,
        "skip_punch_requirement": False,
        "proof_required": False,
        "missing_proof_action": "review",
        "cross_day_allowed": True,
        "deduction_category": "none",
        "enabled": True,
    })
    policy["company"]["overtime"] = {
        "enabled": True,
        "start_mode": "after_normal_end",
        "minimum_minutes": 45,
        "rounding_mode": "actual_minutes",
        "approval_required": True,
        "cross_midnight_allowed": True,
        "cash_enabled": True,
        "cash_formula": "hourly_rate_multiplier",
        "cash_rate": "1",
        "day_rates": {"workday": "1.5", "rest_day": "2", "holiday": "3"},
        "comp_time_enabled": True,
        "comp_time_ratio": "1.5",
    }
    return policy


def _overtime_attendance_workbook(
    workday_type: str | None = None,
    *,
    attendance_text: str = "正常(09:30),正常(20:00);加班(18:00-20:00)",
) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    headers = ["工号", "日期", "每日考勤结果"]
    if workday_type is not None:
        headers.append("工作日类型")
    sheet.append(headers)
    row = [
        "SYN001", "2026-08-01",
        attendance_text,
    ]
    if workday_type is not None:
        row.append(workday_type)
    sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


if __name__ == "__main__":
    unittest.main()
