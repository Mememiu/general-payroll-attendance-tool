from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from app import create_app
from payroll_tool.config.backup import (
    BackupConflictError,
    export_backup,
    import_backup,
)
from payroll_tool.config.legacy_migration import migrate_hr_workbook
from payroll_tool.config.policy_repository import PolicyRepository
from tests.test_app import _completed_hr_workbook
from tests.test_policy_validation import complete_policy


class PolicyBackupTests(unittest.TestCase):
    def test_legacy_workbook_migration_creates_unpublished_draft(self):
        migrated = migrate_hr_workbook(_completed_hr_workbook().getvalue())

        self.assertEqual(migrated["status"], "draft")
        self.assertEqual(migrated["snapshot"]["company"]["worktime"]["daily_minutes"], 480)
        self.assertTrue(migrated["requires_hr_review"])

    def test_legacy_workbook_rejects_work_hours_not_exact_to_a_minute(self):
        from openpyxl import load_workbook

        workbook = load_workbook(_completed_hr_workbook())
        base = workbook["基础与考勤"]
        for row in range(2, base.max_row + 1):
            if base.cell(row, 1).value == "daily_required_hours":
                base.cell(row, 3).value = 7.3333
        content = BytesIO()
        workbook.save(content)
        workbook.close()

        with self.assertRaisesRegex(ValueError, "整分钟"):
            migrate_hr_workbook(content.getvalue())

    def test_backup_round_trip_preserves_version_identity_and_digest(self):
        with TemporaryDirectory() as source_dir, TemporaryDirectory() as target_dir:
            source = PolicyRepository(Path(source_dir) / "policies.sqlite3")
            draft = source.create_draft("公司制度", "2026.01", "2026-01-01")
            source.save_draft(draft.id, complete_policy())
            published = source.publish(draft.id)

            content = export_backup(source)
            target = PolicyRepository(Path(target_dir) / "policies.sqlite3")
            result = import_backup(target, content)
            restored = target.get(published.id)

            self.assertEqual(result.imported, 1)
            self.assertEqual(restored.snapshot_sha256, published.snapshot_sha256)
            self.assertEqual(restored.status.value, "published")

    def test_same_version_with_different_digest_is_rejected(self):
        with TemporaryDirectory() as source_dir, TemporaryDirectory() as target_dir:
            source = PolicyRepository(Path(source_dir) / "policies.sqlite3")
            draft = source.create_draft("公司制度", "2026.01", "2026-01-01")
            source.save_draft(draft.id, complete_policy())
            source.publish(draft.id)
            backup = export_backup(source)

            target = PolicyRepository(Path(target_dir) / "policies.sqlite3")
            import_backup(target, backup)
            changed = _mutate_backup_snapshot(backup)

            with self.assertRaises(BackupConflictError):
                import_backup(target, changed)

    def test_backup_rejects_unknown_format_and_hash_tampering(self):
        with TemporaryDirectory() as directory:
            repository = PolicyRepository(Path(directory) / "policies.sqlite3")
            with self.assertRaisesRegex(ValueError, "format"):
                import_backup(repository, _archive({"format_version": 999}, []))

            valid = export_backup(repository)
            with self.assertRaisesRegex(ValueError, "integrity"):
                import_backup(repository, _replace_policies_without_manifest_hash(valid))

    def test_conflict_rolls_back_every_version_in_the_import(self):
        with TemporaryDirectory() as source_dir, TemporaryDirectory() as target_dir:
            source = PolicyRepository(Path(source_dir) / "policies.sqlite3")
            conflicting = source.create_draft("冲突制度", "1", "2026-01-01")
            conflicting = source.save_draft(conflicting.id, complete_policy())
            first = source.create_draft("新制度", "1", "2026-01-01")
            first = source.save_draft(first.id, complete_policy())
            content = export_backup(source)

            target = PolicyRepository(Path(target_dir) / "policies.sqlite3")
            target.restore_version(conflicting.as_dict())
            changed = _mutate_named_policy(content, "冲突制度")

            with self.assertRaises(BackupConflictError):
                import_backup(target, changed)
            with self.assertRaisesRegex(ValueError, "does not exist"):
                target.get(first.id)

    def test_same_digest_with_changed_metadata_is_a_conflict(self):
        with TemporaryDirectory() as source_dir, TemporaryDirectory() as target_dir:
            source = PolicyRepository(Path(source_dir) / "policies.sqlite3")
            draft = source.create_draft("公司制度", "1", "2026-01-01")
            content = export_backup(source)
            target = PolicyRepository(Path(target_dir) / "policies.sqlite3")
            import_backup(target, content)
            manifest, policies = _read_archive(content)
            policies[0]["effective_from"] = "2026-02-01"

            with self.assertRaises(BackupConflictError):
                import_backup(target, _archive(manifest, policies, refresh_hash=True))


class PolicyBackupApiTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.app = create_app({
            "TESTING": True,
            "POLICY_DATABASE": Path(self.directory.name) / "policies.sqlite3",
        })
        self.client = self.app.test_client()

    def tearDown(self):
        self.app.extensions["task_store"].close()
        self.directory.cleanup()

    def test_export_and_import_endpoints_round_trip(self):
        repository = self.app.extensions["policy_repository"]
        repository.create_draft("公司制度", "2026.01", "2026-01-01")

        exported = self.client.get("/api/advanced/policy-backup")
        self.assertEqual(exported.status_code, 200)
        self.assertIn("attachment", exported.headers["Content-Disposition"])

        imported = self.client.post(
            "/api/advanced/policy-backup/import",
            data={"backup": (BytesIO(exported.data), "policy-backup.zip")},
        )
        self.assertEqual(imported.status_code, 200)
        self.assertEqual(imported.get_json()["skipped"], 1)

    def test_legacy_migration_endpoint_saves_reviewable_draft(self):
        response = self.client.post(
            "/api/advanced/migrate-legacy",
            data={
                "name": "旧制度迁移",
                "version": "legacy-2026",
                "effective_from": "2026-01-01",
                "workbook": (_completed_hr_workbook(), "rules.xlsx"),
            },
        )
        self.assertEqual(response.status_code, 201)
        item = response.get_json()["item"]
        self.assertEqual(item["status"], "draft")
        self.assertTrue(item["snapshot"]["legacy_migration"]["requires_hr_review"])


def _mutate_backup_snapshot(content: bytes) -> bytes:
    from payroll_tool.config.policy_serialization import canonical_snapshot

    manifest, policies = _read_archive(content)
    policies[0]["snapshot"]["company"]["worktime"]["daily_minutes"] = 600
    policies[0]["snapshot_sha256"] = canonical_snapshot(policies[0]["snapshot"]).sha256
    return _archive(manifest, policies, refresh_hash=True)


def _replace_policies_without_manifest_hash(content: bytes) -> bytes:
    manifest, policies = _read_archive(content)
    policies.append({"tampered": True})
    return _archive(manifest, policies)


def _mutate_named_policy(content: bytes, name: str) -> bytes:
    from payroll_tool.config.policy_serialization import canonical_snapshot

    manifest, policies = _read_archive(content)
    item = next(item for item in policies if item["name"] == name)
    item["snapshot"]["company"]["worktime"]["daily_minutes"] = 600
    item["snapshot_sha256"] = canonical_snapshot(item["snapshot"]).sha256
    return _archive(manifest, policies, refresh_hash=True)


def _read_archive(content: bytes):
    with ZipFile(BytesIO(content)) as archive:
        return (
            json.loads(archive.read("manifest.json")),
            json.loads(archive.read("policies.json")),
        )


def _archive(manifest: dict, policies: list, *, refresh_hash: bool = False) -> bytes:
    policies_bytes = json.dumps(
        policies, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    if refresh_hash:
        from hashlib import sha256
        manifest["policies_sha256"] = sha256(policies_bytes).hexdigest()
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("manifest.json", json.dumps(manifest).encode("utf-8"))
        archive.writestr("policies.json", policies_bytes)
    return output.getvalue()


if __name__ == "__main__":
    unittest.main()
