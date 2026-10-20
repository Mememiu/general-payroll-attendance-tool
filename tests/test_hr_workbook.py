from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook

from payroll_tool.templates.hr_workbook import (
    HR_WORKBOOK_SHEETS,
    build_hr_workbook_template,
    validate_hr_workbook,
)


class HrWorkbookTests(unittest.TestCase):
    def _workbook_with_base_values(self, values):
        workbook = load_workbook(BytesIO(build_hr_workbook_template()))
        base = workbook["基础与考勤"]
        rows = {
            base.cell(row, 1).value: row
            for row in range(2, base.max_row + 1)
        }
        for code, value in values.items():
            base.cell(rows[code], 3).value = value
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    def test_template_contains_complete_fixed_sheet_structure(self):
        content = build_hr_workbook_template()
        workbook = load_workbook(BytesIO(content), data_only=False)
        try:
            self.assertEqual(workbook.sheetnames, list(HR_WORKBOOK_SHEETS))
            self.assertEqual(workbook["基础与考勤"]["C2"].value, "1.2.0")
            self.assertGreaterEqual(workbook["假勤与豁免"].max_row, 10)
            self.assertEqual(workbook["五险一金"]["D2"].value, None)
            self.assertIn("employee_id", [cell.value for cell in workbook["员工薪酬"][1]])
            self.assertEqual(workbook["员工薪酬"]["B1"].value, "statutory_enabled")
            self.assertGreater(len(workbook["基础与考勤"].data_validations.dataValidation), 0)
            base_codes = {workbook["基础与考勤"].cell(row, 1).value for row in range(2, workbook["基础与考勤"].max_row + 1)}
            self.assertIn("allowance_taxable", base_codes)
            self.assertIn("bonus_included_in_fund_base", base_codes)
            self.assertIn("checkout_boundary_mode", base_codes)
            self.assertIn("flexible_rule_mode", base_codes)
            self.assertIn("flexible_latest_start", base_codes)
            self.assertIn("core_start", base_codes)
            self.assertIn("core_end", base_codes)
            self.assertIn("overtime_start_mode", base_codes)
        finally:
            workbook.close()

    def test_attendance_strategy_fields_are_blank_and_have_dropdowns(self):
        workbook = load_workbook(BytesIO(build_hr_workbook_template()))
        try:
            base = workbook["基础与考勤"]
            rows = {
                base.cell(row, 1).value: row
                for row in range(2, base.max_row + 1)
            }
            for code in (
                "flexible_rule_mode", "flexible_latest_start", "core_start",
                "core_end", "overtime_start_mode",
            ):
                self.assertIsNone(base.cell(rows[code], 3).value)

            validations = {
                str(validation.sqref): validation.formula1
                for validation in base.data_validations.dataValidation
            }
            self.assertEqual(
                validations[f"C{rows['flexible_rule_mode']}"],
                '"纯工时弹性,最晚到岗限制,核心时段覆盖"',
            )
            self.assertEqual(
                validations[f"C{rows['overtime_start_mode']}"],
                '"审批开始时间,正常下班后,审批开始与正常下班取较晚"',
            )
        finally:
            workbook.close()

    def test_instruction_sheet_matches_template_version(self):
        workbook = load_workbook(BytesIO(build_hr_workbook_template()))
        try:
            self.assertEqual(workbook["填写说明"]["B1"].value, "版本 1.2.0")
        finally:
            workbook.close()

    def test_daily_work_duration_help_does_not_suggest_a_standard_value(self):
        workbook = load_workbook(BytesIO(build_hr_workbook_template()))
        try:
            base = workbook["基础与考勤"]
            rows = {
                base.cell(row, 1).value: row
                for row in range(2, base.max_row + 1)
            }
            help_text = base.cell(rows["daily_required_hours"], 5).value
            self.assertEqual(help_text, "由 HR 按公司制度填写，支持精确到分钟")
        finally:
            workbook.close()

    def test_generated_template_passes_structural_validation(self):
        report = validate_hr_workbook(build_hr_workbook_template())

        self.assertTrue(report.ok)
        self.assertEqual(report.template_version, "1.2.0")

    def test_workbook_missing_required_sheet_is_rejected(self):
        workbook = Workbook()
        content = BytesIO()
        workbook.save(content)
        workbook.close()

        report = validate_hr_workbook(content.getvalue())

        self.assertFalse(report.ok)
        self.assertTrue(any("缺少工作表" in error for error in report.errors))

    def test_blank_template_is_not_complete_enough_for_calculation(self):
        report = validate_hr_workbook(build_hr_workbook_template(), require_complete=True)

        self.assertFalse(report.ok)
        self.assertTrue(any("基础与考勤" in error and "必填" in error for error in report.errors))
        self.assertTrue(any("五险一金" in error for error in report.errors))

    def test_complete_validation_rejects_invalid_rates(self):
        content = BytesIO(build_hr_workbook_template())
        workbook = load_workbook(content)
        workbook["假勤与豁免"]["C2"] = 1.5
        workbook["五险一金"]["D2"] = "是"
        workbook["五险一金"]["L2"] = 1.2
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        report = validate_hr_workbook(output.getvalue(), require_complete=True)

        self.assertTrue(any("计薪比例" in error for error in report.errors))
        self.assertTrue(any("个人费率" in error for error in report.errors))

    def test_complete_validation_rejects_invalid_employee_base_override(self):
        content = BytesIO(build_hr_workbook_template())
        workbook = load_workbook(content)
        employee = workbook["员工薪酬"]
        employee.append(["SYN001", "否", "Synthetic", "正职", None, None, None, 10000, None, None, 0, 0, None, None, None, None, "错误模式", None])
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        report = validate_hr_workbook(output.getvalue(), require_complete=True)

        self.assertTrue(any("social_base_mode" in error for error in report.errors))

    def test_fixed_mode_must_not_carry_flexible_strategy(self):
        report = validate_hr_workbook(
            self._workbook_with_base_values({
                "work_mode": "固定工时",
                "flexible_rule_mode": "纯工时弹性",
            }),
            require_complete=True,
        )

        self.assertTrue(any("固定工时" in error and "flexible_rule_mode" in error for error in report.errors))

    def test_work_mode_must_be_a_supported_value(self):
        report = validate_hr_workbook(
            self._workbook_with_base_values({"work_mode": "弹性"}),
            require_complete=True,
        )

        self.assertTrue(any("work_mode" in error for error in report.errors))

    def test_overtime_start_mode_must_be_a_supported_value(self):
        report = validate_hr_workbook(
            self._workbook_with_base_values({"overtime_start_mode": "自动判断"}),
            require_complete=True,
        )

        self.assertTrue(any("overtime_start_mode" in error for error in report.errors))

    def test_flexible_mode_requires_supported_strategy(self):
        for strategy in (None, "任意策略"):
            with self.subTest(strategy=strategy):
                report = validate_hr_workbook(
                    self._workbook_with_base_values({
                        "work_mode": "弹性工时",
                        "flexible_rule_mode": strategy,
                    }),
                    require_complete=True,
                )

                self.assertTrue(any("flexible_rule_mode" in error for error in report.errors))

    def test_latest_start_strategy_requires_flexible_latest_start(self):
        report = validate_hr_workbook(
            self._workbook_with_base_values({
                "work_mode": "弹性工时",
                "flexible_rule_mode": "最晚到岗限制",
            }),
            require_complete=True,
        )

        self.assertTrue(any("flexible_latest_start" in error for error in report.errors))

    def test_core_hours_strategy_requires_both_core_times(self):
        for missing_code in ("core_start", "core_end"):
            values = {
                "work_mode": "弹性工时",
                "flexible_rule_mode": "核心时段覆盖",
                "core_start": "10:00",
                "core_end": "16:00",
            }
            values[missing_code] = None
            with self.subTest(missing_code=missing_code):
                report = validate_hr_workbook(
                    self._workbook_with_base_values(values),
                    require_complete=True,
                )

                self.assertTrue(any(missing_code in error for error in report.errors))

    def test_attendance_strategy_times_must_be_hh_mm_when_supplied(self):
        for code in ("flexible_latest_start", "core_start", "core_end", "grace_checkout"):
            with self.subTest(code=code):
                report = validate_hr_workbook(
                    self._workbook_with_base_values({code: "9点30"}),
                    require_complete=True,
                )

                self.assertTrue(any(code in error and "HH:MM" in error for error in report.errors))

    def test_blank_grace_checkout_does_not_add_a_completeness_error(self):
        report = validate_hr_workbook(
            self._workbook_with_base_values({"grace_checkout": None}),
            require_complete=True,
        )

        self.assertFalse(any("grace_checkout" in error for error in report.errors))


if __name__ == "__main__":
    unittest.main()
