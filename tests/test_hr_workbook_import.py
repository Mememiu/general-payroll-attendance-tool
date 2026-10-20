from datetime import date
from decimal import Decimal
from io import BytesIO
import unittest

from openpyxl import load_workbook

from payroll_tool.importers.hr_workbook import attendance_policy_from_hr, income_item_policies_from_hr, parse_hr_workbook
from payroll_tool.templates.hr_workbook import build_hr_workbook_template


class HrWorkbookImportTests(unittest.TestCase):
    def test_parser_reads_explicit_attendance_strategies(self):
        policy = attendance_policy_from_hr(parse_hr_workbook(_completed_workbook()))

        self.assertEqual(policy.overtime_start_mode, "审批开始时间")
        self.assertIsNone(policy.flexible_rule_mode)

    def test_parser_reads_core_time_flexible_policy(self):
        workbook = load_workbook(BytesIO(_completed_workbook()))
        _set_base_values(workbook, {
            "work_mode": "弹性工时",
            "flexible_rule_mode": "核心时段覆盖",
            "core_start": "10:30",
            "core_end": "16:30",
        })
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        policy = attendance_policy_from_hr(parse_hr_workbook(output.getvalue()))

        self.assertEqual(policy.flexible_rule_mode, "核心时段覆盖")
        self.assertEqual(policy.core_start, "10:30")
        self.assertEqual(policy.core_end, "16:30")

    def test_parser_rejects_daily_work_duration_with_fractional_minute(self):
        workbook = load_workbook(BytesIO(_completed_workbook()))
        _set_base_values(workbook, {"daily_required_hours": 6.81})
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        with self.assertRaisesRegex(ValueError, "精确到整分钟"):
            attendance_policy_from_hr(parse_hr_workbook(output.getvalue()))

    def test_parser_returns_mapping_and_employee_profiles_from_fixed_workbook(self):
        content = _completed_workbook()

        result = parse_hr_workbook(content)

        self.assertEqual(result.attendance_mapping["employee_id"], "工号")
        self.assertEqual(result.base_values["work_mode"], "固定工时")
        self.assertEqual(len(result.employees), 1)
        employee = result.employees[0]
        self.assertEqual(employee.employee_id, "SYN001")
        self.assertEqual(employee.employee_name, "Synthetic")
        self.assertIsNone(employee.hire_date)
        self.assertIsNone(employee.termination_date)
        self.assertEqual(employee.employment_type, "full_time")
        self.assertEqual(employee.monthly_salary, Decimal("10000"))
        self.assertFalse(employee.statutory_enabled)
        self.assertEqual(result.leave_rules["病假"].paid_ratio, Decimal("1"))
        self.assertTrue(result.leave_rules["病假"].counts_as_absence)
        policy = attendance_policy_from_hr(result)
        self.assertEqual(policy.standard_start, "09:00")
        self.assertEqual(policy.daily_required_hours, 8)
        income = income_item_policies_from_hr(result)
        self.assertTrue(income["allowance"].taxable)
        self.assertFalse(income["bonus"].included_in_fund_base)

    def test_parser_reads_configurable_attendance_boundaries(self):
        workbook = load_workbook(BytesIO(_completed_workbook()))
        base = workbook["基础与考勤"]
        values = {"grace_checkout": "20:00", "full_day_absence_after_hours": 2.5}
        for row in range(2, base.max_row + 1):
            if base.cell(row, 1).value in values:
                base.cell(row, 3).value = values[base.cell(row, 1).value]
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        policy = attendance_policy_from_hr(parse_hr_workbook(output.getvalue()))

        self.assertEqual(policy.grace_checkout, "20:00")
        self.assertEqual(policy.full_day_absence_after_hours, 2.5)

    def test_parser_reads_enabled_statutory_policy_rows(self):
        workbook = load_workbook(BytesIO(_completed_workbook()))
        social = workbook["五险一金"]
        social["C2"], social["D2"], social["E2"] = "standard", "是", "Synthetic City"
        social["F2"], social["H2"] = "2026-01-01", "薪资全额"
        social["I2"], social["J2"], social["L2"], social["M2"] = 5000, 20000, 0.08, 0.16
        social["N2"], social["O2"], social["P2"] = "Synthetic Notice", "https://example.test", "2026-01-01"
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        result = parse_hr_workbook(output.getvalue())

        rule = result.statutory_rules[0]
        self.assertEqual(rule.item_code, "pension")
        self.assertEqual(rule.employee_rate, Decimal("0.08"))
        self.assertEqual(rule.base_mode, "薪资全额")

    def test_parser_reads_employee_cumulative_tax_record(self):
        workbook = load_workbook(BytesIO(_completed_workbook()))
        workbook["个税累计"].append(["SYN001", 2026, 30000, 1200, 500, "synthetic"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        result = parse_hr_workbook(output.getvalue())

        record = result.cumulative_tax["SYN001"]
        self.assertEqual(record.tax_year, 2026)
        self.assertEqual(record.prior_withheld_tax, Decimal("500"))

    def test_parser_reads_employee_statutory_overrides(self):
        workbook = load_workbook(BytesIO(_completed_workbook()))
        employee = workbook["员工薪酬"]
        row = next(
            row for row in range(2, employee.max_row + 1)
            if employee.cell(row, 1).value == "SYN001"
        )
        employee.cell(row, 14).value = "special"
        employee.cell(row, 17).value = "固定基数"
        employee.cell(row, 18).value = 6000
        employee.cell(row, 19).value = "薪资全额"
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        result = parse_hr_workbook(output.getvalue())

        parsed = result.employees[0]
        self.assertEqual(parsed.pension_tier, "special")
        self.assertEqual(parsed.social_base_mode, "固定基数")
        self.assertEqual(parsed.social_fixed_base, Decimal("6000"))
        self.assertEqual(parsed.fund_base_mode, "薪资全额")

    def test_parser_reads_employee_perfect_attendance_override(self):
        workbook = load_workbook(BytesIO(_completed_workbook()))
        employee = workbook["员工薪酬"]
        row = next(
            row for row in range(2, employee.max_row + 1)
            if employee.cell(row, 1).value == "SYN001"
        )
        employee.cell(row, 13).value = 150
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        parsed = parse_hr_workbook(output.getvalue()).employees[0]

        self.assertEqual(parsed.perfect_attendance_override, Decimal("150"))


def _completed_workbook() -> bytes:
    workbook = load_workbook(BytesIO(build_hr_workbook_template()))
    base_values = {
        "package_id": "synthetic-policy", "policy_version": "1", "work_mode": "固定工时",
        "overtime_start_mode": "审批开始时间",
        "daily_required_hours": 8, "monthly_pay_days": 21.75, "standard_start": "09:00",
        "standard_end": "18:30", "lunch_start": "12:00", "lunch_end": "13:30",
        "earliest_valid_start": "08:00", "latest_normal_start": "09:15", "late_grace_end": "09:30",
        "late_makeup_allowed": "否", "full_day_absence_after_hours": 2,
        "attendance_deduction_factor": 1, "sick_leave_deduction_factor": 0.3,
        "perfect_attendance_award": 0, "overtime_cash_enabled": "否",
        "entry_day_full_attendance": "是", "termination_day_full_attendance": "是",
        "allowance_taxable": "是", "allowance_included_in_social_base": "否",
        "allowance_included_in_fund_base": "否", "bonus_taxable": "是",
        "bonus_included_in_social_base": "否", "bonus_included_in_fund_base": "否",
    }
    base = workbook["基础与考勤"]
    for row in range(2, base.max_row + 1):
        if base.cell(row, 1).value in base_values:
            base.cell(row, 3).value = base_values[base.cell(row, 1).value]
    mapping = workbook["考勤字段映射"]
    mapping["D2"], mapping["D3"], mapping["D4"] = "工号", "日期", "每日考勤结果"
    leave = workbook["假勤与豁免"]
    for row in range(2, leave.max_row + 1):
        for col, value in {3: 1, 4: "否", 5: "是", 6: "是", 7: "是", 8: "否"}.items():
            leave.cell(row, col).value = value
    social = workbook["五险一金"]
    for row in range(2, social.max_row + 1):
        social.cell(row, 4).value = "否"
    workbook["员工薪酬"].append(["SYN001", "否", "Synthetic", "正职", None, None, None, 10000, None, None, 0, 0])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _set_base_values(workbook, values: dict[str, object]) -> None:
    base = workbook["基础与考勤"]
    for row in range(2, base.max_row + 1):
        code = base.cell(row, 1).value
        if code in values:
            base.cell(row, 3).value = values[code]
