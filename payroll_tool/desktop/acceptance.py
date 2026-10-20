"""Privacy-safe evidence for desktop end-to-end acceptance."""

from hashlib import sha256
from io import BytesIO

from openpyxl import load_workbook


def workbook_evidence(content):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheet_names = workbook.sheetnames
        results = workbook["计算结果"]
        rows = results.iter_rows()
        headers = [cell.value for cell in next(rows)]
        employee_index = headers.index("employee_id")
        employee_ids = sorted(str(row[employee_index].value) for row in rows if row[employee_index].value is not None)
        return {
            "result_sha256": sha256(content).hexdigest(),
            "sheet_names": sheet_names,
            "employee_count": len(employee_ids),
            "employee_ids_sha256": sha256("\n".join(employee_ids).encode()).hexdigest(),
        }
    finally:
        workbook.close()
