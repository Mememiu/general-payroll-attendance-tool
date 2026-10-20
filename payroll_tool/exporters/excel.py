"""Standard payroll workbook export with privacy-safe audit metadata."""

from io import BytesIO
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from payroll_tool.audit import AuditManifest


def build_payroll_export(
    rows: list[dict[str, object]], audit: AuditManifest
) -> bytes:
    workbook = Workbook()
    results = workbook.active
    results.title = "计算结果"
    headers = list(rows[0]) if rows else ["employee_id"]
    results.append(headers)
    for row in rows:
        results.append([_excel_value(row.get(header)) for header in headers])
    _format_sheet(results)

    audit_sheet = workbook.create_sheet("审计信息")
    audit_sheet.append(["字段", "值"])
    for field, value in audit.to_dict().items():
        audit_sheet.append([field, value])
    _format_sheet(audit_sheet)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _excel_value(value: object) -> object:
    if isinstance(value, Decimal):
        return float(value)
    return value


def _format_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="17324D")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(32, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
