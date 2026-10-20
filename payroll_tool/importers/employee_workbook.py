"""Import the fixed employee payroll workbook independently from HR policies."""

from collections.abc import Callable, Mapping, Set
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from payroll_tool.payroll.engine import EmployeePayProfile
from payroll_tool.templates.hr_workbook import EXPECTED_HEADERS


SHEET_NAME = "员工薪酬"
REQUIRED_COLUMNS = {"employee_id", "employment_type"}
_ZERO_ID_FORMAT = re.compile(r"^0+$")


class FieldValueError(ValueError):
    def __init__(self, field: str, reason: str):
        self.field = field
        super().__init__(reason)


def parse_employee_workbook(
    content: bytes,
    allowed_employment_types: Set[str] | Mapping[str, object],
    statutory_tier_resolver: (
        Callable[[EmployeePayProfile], Mapping[str, Set[str]] | None] | None
    ) = None,
    *,
    employee_transform: (
        Callable[[EmployeePayProfile], EmployeePayProfile] | None
    ) = None,
) -> tuple[EmployeePayProfile, ...]:
    if not isinstance(content, bytes) or not content:
        raise ValueError("员工薪酬文件为空或无效")
    policy_by_type = (
        allowed_employment_types if isinstance(allowed_employment_types, Mapping) else {}
    )
    allowed = {
        value.strip()
        for value in allowed_employment_types
        if isinstance(value, str) and value.strip()
    }
    if not allowed:
        raise ValueError("所选规则未配置可用的用工类型")
    try:
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ValueError("员工薪酬文件不是有效的 XLSX") from exc
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError("员工薪酬文件缺少“员工薪酬”工作表")
        sheet = workbook[SHEET_NAME]
        headers = [_clean_text(cell.value) for cell in sheet[1]]
        if len(headers) != len(set(headers)):
            raise ValueError("员工薪酬表存在重复列名")
        missing = REQUIRED_COLUMNS - set(headers)
        if missing:
            raise ValueError(f"员工薪酬表缺少必填列：{','.join(sorted(missing))}")

        employees: list[EmployeePayProfile] = []
        seen_ids: set[str] = set()
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            if not any(cell.value is not None and str(cell.value).strip() for cell in cells):
                continue
            row = {header: cell.value for header, cell in zip(headers, cells, strict=True)}
            try:
                employee_id = _employee_id(cells[headers.index("employee_id")])
            except ValueError as exc:
                raise ValueError(
                    f"员工薪酬第 {row_number} 行字段 employee_id: {exc}"
                ) from exc
            if not employee_id:
                raise ValueError(f"员工薪酬第 {row_number} 行缺少工号")
            if employee_id in seen_ids:
                raise ValueError(f"员工薪酬工号重复：{employee_id}")
            seen_ids.add(employee_id)
            employment_type = _clean_text(row.get("employment_type"))
            if employment_type not in allowed:
                raise ValueError(
                    f"员工薪酬第 {row_number} 行（工号 {employee_id}）"
                    f"字段 employment_type: 用工类型 {employment_type or '（空）'} "
                    "不在当前已发布规则中"
                )
            row["employee_id"] = employee_id
            try:
                _validate_employee_options(row)
                _validate_non_negative_amounts(row)
                employee = employee_from_row(row)
                if employee_transform is not None:
                    employee = employee_transform(employee)
                if employee.employment_type in policy_by_type:
                    _validate_payroll_scheme_fields(
                        employee,
                        policy_by_type[employee.employment_type].payroll_scheme,
                    )
                if statutory_tier_resolver is not None:
                    allowed_tiers = statutory_tier_resolver(employee)
                    if allowed_tiers is not None:
                        _validate_statutory_tiers(employee, allowed_tiers)
                employees.append(employee)
            except FieldValueError as exc:
                raise ValueError(
                    f"员工薪酬第 {row_number} 行（工号 {employee_id}）"
                    f"字段 {exc.field}: {exc}"
                ) from exc
        if not employees:
            raise ValueError("员工薪酬表至少需要填写一名员工")
        return tuple(employees)
    finally:
        workbook.close()


def employee_from_row(
    row: Mapping[str, object],
    employment_type_map: Mapping[str, str] | None = None,
) -> EmployeePayProfile:
    source_type = _clean_text(row.get("employment_type"))
    employment_type = (
        employment_type_map.get(source_type, source_type)
        if employment_type_map is not None
        else source_type
    )
    return EmployeePayProfile(
        employee_id=_clean_text(row.get("employee_id")),
        employment_type=employment_type,
        monthly_salary=_field_value(row, "monthly_salary", decimal_or_none),
        intern_daily_rate=_field_value(row, "intern_daily_rate", decimal_or_none),
        intern_hourly_rate=_field_value(row, "intern_hourly_rate", decimal_or_none),
        conversion_date=_field_value(row, "conversion_date", date_or_none),
        allowances=_field_value(row, "allowance", decimal_or_zero),
        bonuses=_field_value(row, "bonus", decimal_or_zero),
        perfect_attendance_override=_field_value(
            row, "perfect_attendance_override", decimal_or_none
        ),
        employee_name=text_or_none(row.get("employee_name")) or "",
        hire_date=_field_value(row, "hire_date", date_or_none),
        termination_date=_field_value(row, "termination_date", date_or_none),
        statutory_enabled=yes(row.get("statutory_enabled")),
        pension_tier=text_or_none(row.get("pension_tier")),
        medical_tier=text_or_none(row.get("medical_tier")),
        work_injury_tier=text_or_none(row.get("work_injury_tier")),
        social_base_mode=text_or_none(row.get("social_base_mode")),
        social_fixed_base=_field_value(row, "social_fixed_base", decimal_or_none),
        fund_base_mode=text_or_none(row.get("fund_base_mode")),
        fund_fixed_base=_field_value(row, "fund_fixed_base", decimal_or_none),
        tax_year=_field_value(row, "tax_year", integer_or_none),
        prior_taxable_income=_field_value(row, "prior_taxable_income", decimal_or_zero),
        prior_special_deductions=_field_value(
            row, "prior_special_deductions", decimal_or_zero
        ),
        current_special_deductions=_field_value(
            row, "current_special_deductions", decimal_or_zero
        ),
        prior_withheld_tax=_field_value(row, "prior_withheld_tax", decimal_or_zero),
    )


def _field_value(row: Mapping[str, object], field: str, converter):
    try:
        return converter(row.get(field))
    except ValueError as exc:
        raise FieldValueError(field, str(exc)) from exc


def _validate_employee_options(row: Mapping[str, object]) -> None:
    _field_value(row, "statutory_enabled", yes_or_no)
    _field_value(row, "social_base_mode", base_mode_or_none)
    _field_value(row, "fund_base_mode", base_mode_or_none)
    for mode_field, amount_field in (
        ("social_base_mode", "social_fixed_base"),
        ("fund_base_mode", "fund_fixed_base"),
    ):
        if row.get(mode_field) == "固定基数" and row.get(amount_field) in (None, ""):
            raise FieldValueError(amount_field, f"{mode_field} 为固定基数时该字段必填")


_AMOUNT_FIELDS = (
    "monthly_salary", "intern_daily_rate", "intern_hourly_rate", "allowance",
    "bonus", "perfect_attendance_override", "social_fixed_base",
    "fund_fixed_base", "prior_taxable_income", "prior_special_deductions",
    "current_special_deductions", "prior_withheld_tax",
)


def _validate_non_negative_amounts(row: Mapping[str, object]) -> None:
    for field in _AMOUNT_FIELDS:
        value = _field_value(row, field, decimal_or_none)
        if value is not None and value < 0:
            raise FieldValueError(field, "金额不能为负数")


def _validate_payroll_scheme_fields(employee: EmployeePayProfile, scheme: str) -> None:
    required_by_scheme = {
        "full_time_monthly": ("monthly_salary",),
        "intern_daily_hourly": ("intern_daily_rate", "intern_hourly_rate"),
        "conversion_split": (
            "monthly_salary", "intern_daily_rate", "intern_hourly_rate",
            "conversion_date",
        ),
    }
    for field in required_by_scheme.get(scheme, ()):
        if getattr(employee, field) is None:
            raise FieldValueError(field, f"薪资方案 {scheme} 要求该字段必填")


def _validate_statutory_tiers(
    employee: EmployeePayProfile, allowed_tiers: Mapping[str, Set[str]]
) -> None:
    if not employee.statutory_enabled:
        return
    for item_code, field in {
        "pension": "pension_tier",
        "medical": "medical_tier",
        "work_injury": "work_injury_tier",
    }.items():
        tiers = set(allowed_tiers.get(item_code, ()))
        value = getattr(employee, field)
        if value is not None and value not in tiers:
            raise FieldValueError(field, f"档位代码 {value} 不在已发布规则中")
        if employee.statutory_enabled and len(tiers) > 1 and value is None:
            raise FieldValueError(field, "已发布规则存在多个档位，该字段必填")


def decimal_or_zero(value: object) -> Decimal:
    return decimal_or_none(value) or Decimal("0")


def integer_or_none(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"年份不是有效整数：{value}")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"年份不是有效整数：{value}") from exc
    if not result.is_finite() or result != result.to_integral_value():
        raise ValueError(f"年份不是有效整数：{value}")
    return int(result)


def decimal_or_none(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"金额不是有效数字：{value}") from exc
    if not result.is_finite():
        raise ValueError(f"金额不是有限数字：{value}")
    return result


def date_or_none(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value))
    except ValueError as exc:
        raise ValueError(f"日期必须使用 YYYY-MM-DD：{value}") from exc


def yes_or_no(value: object) -> bool:
    if value is None or value == "":
        raise ValueError("该字段必填，必须填写是或否")
    if value == "否" or value is False:
        return False
    if value == "是" or value is True:
        return True
    raise ValueError(f"必须填写是或否：{value}")


def yes(value: object) -> bool:
    return value == "是" or value is True


def base_mode_or_none(value: object) -> str | None:
    result = text_or_none(value)
    if result is None or result in {"公司默认", "薪资全额", "固定基数"}:
        return result
    raise ValueError(f"必须填写公司默认、薪资全额或固定基数：{value}")


def text_or_none(value: object) -> str | None:
    text = _clean_text(value)
    return text or None


def _employee_id(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        raise ValueError("员工工号不能是布尔值")
    if isinstance(value, int) and _ZERO_ID_FORMAT.fullmatch(cell.number_format or ""):
        return f"{value:0{len(cell.number_format)}d}"
    if isinstance(value, float) and value.is_integer():
        integer = int(value)
        if _ZERO_ID_FORMAT.fullmatch(cell.number_format or ""):
            return f"{integer:0{len(cell.number_format)}d}"
        return str(integer)
    return _clean_text(value)


def _clean_text(value: object) -> str:
    return str(value).strip() if value is not None else ""
