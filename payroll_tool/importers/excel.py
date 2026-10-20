"""Explicit header mapping validation for current-task workbook imports."""

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO

from openpyxl import load_workbook

from payroll_tool.importers.uploaded_inputs import UploadedFile


ATTENDANCE_REQUIRED_FIELDS = {"employee_id", "work_date", "attendance_text"}


@dataclass(frozen=True)
class MappingValidationReport:
    errors: tuple[str, ...]

    @property
    def ok(self) -> bool:
        return not self.errors


@dataclass(frozen=True)
class ImportedAttendance:
    rows: tuple[dict[str, object], ...]
    source_sha256: str


def validate_mapping(
    headers: list[str],
    mapping: dict[str, str],
    required_fields: set[str] | None = None,
) -> MappingValidationReport:
    normalized_headers = {header.strip() for header in headers if header.strip()}
    errors: list[str] = []
    mapped_columns: dict[str, str] = {}
    for field in sorted((required_fields or set()) - set(mapping)):
        errors.append(f"required mapped field is missing: {field}")
    for field, source_column in mapping.items():
        source_column = source_column.strip()
        if not source_column:
            errors.append(f"field {field} has no uploaded source column")
            continue
        if source_column not in normalized_headers:
            errors.append(f"uploaded header not found: {source_column} for {field}")
        previous_field = mapped_columns.get(source_column)
        if previous_field is not None:
            errors.append(
                f"uploaded header {source_column} is mapped more than once: "
                f"{previous_field}, {field}"
            )
        else:
            mapped_columns[source_column] = field
    return MappingValidationReport(errors=tuple(errors))


def import_attendance_workbook(
    upload: UploadedFile, mapping: dict[str, str]
) -> ImportedAttendance:
    if not isinstance(upload, UploadedFile):
        raise TypeError("attendance import requires a current-task UploadedFile")
    content = upload.path.read_bytes()
    current_hash = sha256(content).hexdigest()
    if current_hash != upload.sha256:
        raise ValueError("attendance file changed after upload registration")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("uploaded attendance workbook is not a valid .xlsx file") from exc
    try:
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        first_row = next(values, None)
        if first_row is None:
            raise ValueError("uploaded attendance workbook is empty")
        headers = [str(value).strip() if value is not None else "" for value in first_row]
        report = validate_mapping(headers, mapping, ATTENDANCE_REQUIRED_FIELDS)
        if not report.ok:
            raise ValueError("; ".join(report.errors))
        header_indexes = {header: index for index, header in enumerate(headers)}
        rows = tuple(
            {
                field: row[header_indexes[column]]
                if header_indexes[column] < len(row)
                else None
                for field, column in mapping.items()
            }
            for row in values
            if any(value is not None for value in row)
        )
        return ImportedAttendance(rows=rows, source_sha256=current_hash)
    finally:
        workbook.close()
