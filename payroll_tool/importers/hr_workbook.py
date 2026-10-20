"""Convert the fixed HR workbook into engine-ready input models."""

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from payroll_tool.config.schema import AttendancePolicy, IncomeItemPolicy
from payroll_tool.payroll.engine import EmployeePayProfile
from payroll_tool.templates.hr_workbook import validate_hr_workbook
from payroll_tool.importers.employee_workbook import employee_from_row


EMPLOYMENT_TYPES = {"正职": "full_time", "实习": "intern", "实习转正": "converted"}


@dataclass(frozen=True)
class ParsedHrWorkbook:
    base_values: dict[str, object]
    attendance_mapping: dict[str, str]
    employees: tuple[EmployeePayProfile, ...]
    leave_rules: dict[str, "LeaveRuleInput"]
    statutory_rules: tuple["StatutoryRuleInput", ...]
    cumulative_tax: dict[str, "CumulativeTaxInput"]


@dataclass(frozen=True)
class LeaveRuleInput:
    paid_ratio: Decimal
    counts_as_attendance: bool
    counts_as_absence: bool
    affects_perfect_attendance: bool
    punch_exempt: bool
    code: str | None = None
    pay_scheme: str = "ratio"
    fixed_deduction_amount: Decimal | None = None
    formula_id: str | None = None
    formula_multiplier: Decimal | None = None
    deduction_category: str = "none"


@dataclass(frozen=True)
class StatutoryRuleInput:
    item_code: str
    tier_code: str
    region: str
    effective_from: date
    effective_to: date | None
    base_mode: str
    base_lower: Decimal
    base_upper: Decimal
    fixed_base: Decimal | None
    employee_rate: Decimal
    employer_rate: Decimal
    employment_types: tuple[str, ...] = ()
    hire_month_rule: str = "always"
    termination_month_rule: str = "always"
    hire_cutoff_day: int | None = None
    termination_cutoff_day: int | None = None


@dataclass(frozen=True)
class CumulativeTaxInput:
    tax_year: int
    prior_taxable_income: Decimal
    prior_special_deductions: Decimal
    prior_withheld_tax: Decimal


def parse_hr_workbook(content: bytes) -> ParsedHrWorkbook:
    report = validate_hr_workbook(content, require_complete=True)
    if not report.ok:
        raise ValueError("; ".join(report.errors))
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        base_values = {
            row[0]: row[2]
            for row in workbook["基础与考勤"].iter_rows(min_row=2, values_only=True)
            if row[0]
        }
        attendance_mapping = {
            row[0]: str(row[3]).strip()
            for row in workbook["考勤字段映射"].iter_rows(min_row=2, values_only=True)
            if row[0] and row[3]
        }
        employees = tuple(_employee_from_row(row) for row in _data_rows(workbook["员工薪酬"]))
        leave_rules = {
            str(row["规则名称"]): LeaveRuleInput(
                paid_ratio=_decimal_or_zero(row["计薪比例"]),
                counts_as_attendance=_yes(row["计入出勤"]),
                counts_as_absence=_yes(row["计入缺勤"]),
                affects_perfect_attendance=_yes(row["影响全勤"]),
                punch_exempt=_yes(row["免打卡"]),
            )
            for row in _data_rows(workbook["假勤与豁免"])
        }
        statutory_rules = tuple(
            StatutoryRuleInput(
                item_code=str(row["item_code"]),
                tier_code=str(row["档位代码"]),
                region=str(row["地区"]),
                effective_from=_date_or_none(row["生效日"]),
                effective_to=_date_or_none(row["失效日"]),
                base_mode=str(row["基数模式"]),
                base_lower=_decimal_or_zero(row["基数下限"]),
                base_upper=_decimal_or_zero(row["基数上限"]),
                fixed_base=_decimal_or_none(row["固定基数"]),
                employee_rate=_decimal_or_zero(row["个人费率"]),
                employer_rate=_decimal_or_zero(row["单位费率"]),
            )
            for row in _data_rows(workbook["五险一金"])
            if row["启用"] == "是"
        )
        cumulative_tax = {
            str(row["employee_id"]).strip(): CumulativeTaxInput(
                tax_year=int(row["tax_year"]),
                prior_taxable_income=_decimal_or_zero(row["prior_taxable_income"]),
                prior_special_deductions=_decimal_or_zero(row["prior_special_deductions"]),
                prior_withheld_tax=_decimal_or_zero(row["prior_withheld_tax"]),
            )
            for row in _data_rows(workbook["个税累计"])
        }
        return ParsedHrWorkbook(
            base_values, attendance_mapping, employees, leave_rules, statutory_rules, cumulative_tax
        )
    finally:
        workbook.close()


def attendance_policy_from_hr(parsed: ParsedHrWorkbook) -> AttendancePolicy:
    values = parsed.base_values
    work_mode = {"固定工时": "fixed", "弹性工时": "flexible"}.get(str(values["work_mode"]))
    if work_mode is None:
        raise ValueError("HR 工作簿工时模式无效")
    required_minutes = Decimal(str(values["daily_required_hours"])) * Decimal("60")
    if (
        required_minutes <= 0
        or required_minutes != required_minutes.to_integral_value()
    ):
        raise ValueError("每日标准工时必须是大于0且精确到整分钟的数值")
    return AttendancePolicy(
        standard_start=_time_text(values["standard_start"]),
        standard_end=_time_text(values["standard_end"]),
        lunch_start=_time_text(values["lunch_start"]),
        lunch_end=_time_text(values["lunch_end"]),
        latest_normal_start=_time_text(values["latest_normal_start"]),
        late_grace_end=_time_text(values["late_grace_end"]),
        work_mode=work_mode,
        daily_required_hours=float(required_minutes / Decimal("60")),
        leave_types={},
        attendance_exemptions={},
        overtime_start_mode=_text_or_none(values.get("overtime_start_mode")),
        flexible_rule_mode=_text_or_none(values.get("flexible_rule_mode")),
        flexible_latest_start=_optional_time_text(values.get("flexible_latest_start")),
        core_start=_optional_time_text(values.get("core_start")),
        core_end=_optional_time_text(values.get("core_end")),
        checkout_boundary_mode=_text_or_none(values.get("checkout_boundary_mode")),
        earliest_valid_start=_optional_time_text(values.get("earliest_valid_start")),
        grace_checkout=_optional_time_text(values.get("grace_checkout")),
        late_makeup_allowed=_yes(values.get("late_makeup_allowed")),
        full_day_absence_after_hours=float(values["full_day_absence_after_hours"]),
        entry_day_full_attendance=_yes(values["entry_day_full_attendance"]),
        termination_day_full_attendance=_yes(values["termination_day_full_attendance"]),
    )


def income_item_policies_from_hr(
    parsed: ParsedHrWorkbook,
) -> dict[str, IncomeItemPolicy]:
    values = parsed.base_values
    return {
        "allowance": IncomeItemPolicy(
            taxable=_yes(values["allowance_taxable"]),
            included_in_social_base=_yes(values["allowance_included_in_social_base"]),
            included_in_fund_base=_yes(values["allowance_included_in_fund_base"]),
        ),
        "bonus": IncomeItemPolicy(
            taxable=_yes(values["bonus_taxable"]),
            included_in_social_base=_yes(values["bonus_included_in_social_base"]),
            included_in_fund_base=_yes(values["bonus_included_in_fund_base"]),
        ),
    }


def _data_rows(sheet):
    headers = [cell.value for cell in sheet[1]]
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values, strict=True))
        if any(value is not None and str(value).strip() for value in row.values()):
            yield row


def _employee_from_row(row: dict[str, object]) -> EmployeePayProfile:
    return employee_from_row(row, EMPLOYMENT_TYPES)


def _decimal_or_zero(value: object) -> Decimal:
    return _decimal_or_none(value) or Decimal("0")


def _decimal_or_none(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    return Decimal(str(value))


def _date_or_none(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value))


def _yes(value: object) -> bool:
    return value == "是"


def _text_or_none(value: object) -> str | None:
    text = str(value).strip() if value is not None else ""
    return text or None


def _time_text(value: object) -> str:
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if hasattr(value, "strftime"):
        return value.strftime("%H:%M")
    return str(value).strip()


def _optional_time_text(value: object) -> str | None:
    if value is None or value == "":
        return None
    return _time_text(value)
