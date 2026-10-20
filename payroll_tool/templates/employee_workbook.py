"""Blank employee payroll workbook for the in-app policy workflow."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from payroll_tool.templates.hr_workbook import EXPECTED_HEADERS


FIELD_GUIDANCE = (
    ("employee_id", "员工工号", "必填", "所有员工", "文本", "必须与考勤文件工号完全一致；保留前导零。"),
    ("statutory_enabled", "是否缴纳五险一金", "必填", "所有员工", "是/否", "选择是否参与法定缴费，最终适用范围由用工类型规则决定。"),
    ("employee_name", "员工姓名", "可选", "所有员工", "文本", "用于结果展示，不参与工号匹配。"),
    ("employment_type", "用工类型代码", "必填", "所有员工", "规则中的启用代码", "必须填写已发布 HR 规则中启用的用工类型代码。"),
    ("hire_date", "入职日期", "条件必填", "存在入职折算或缴费判断时", "YYYY-MM-DD", "员工实际入职日期。"),
    ("termination_date", "离职日期", "条件必填", "存在离职折算或缴费判断时", "YYYY-MM-DD", "员工实际离职日期；在职员工留空。"),
    ("conversion_date", "实习转正日期", "条件必填", "实习转正用工类型", "YYYY-MM-DD", "转正日开始采用转正后的薪资方案。"),
    ("monthly_salary", "月薪", "条件必填", "按月薪计薪的用工类型", "非负金额", "税前月薪，不含当月单列奖金和补贴。"),
    ("intern_daily_rate", "实习日薪", "条件必填", "按日薪计薪且规则要求时", "非负金额", "实习生满勤日适用的日薪。"),
    ("intern_hourly_rate", "实习时薪", "条件必填", "按时薪计薪且规则要求时", "非负金额", "实习生非满勤时适用的时薪。"),
    ("allowance", "当月补贴", "可选", "所有员工", "非负金额", "当月补贴；留空按 0 处理。"),
    ("bonus", "当月奖金", "可选", "所有员工", "非负金额", "当月奖金；留空按 0 处理。"),
    ("perfect_attendance_override", "全勤奖员工覆盖", "可选", "需要员工级覆盖时", "非负金额", "优先于公司全局金额；填写 0 可取消该员工全勤奖。"),
    ("pension_tier", "养老保险档位", "条件必填", "启用养老保险且存在多个档位时", "规则中的档位代码", "填写已发布规则中的养老保险档位代码。"),
    ("medical_tier", "医疗保险档位", "条件必填", "启用医疗保险且存在多个档位时", "规则中的档位代码", "填写已发布规则中的医疗保险档位代码。"),
    ("work_injury_tier", "工伤保险档位", "条件必填", "启用工伤保险且存在多个档位时", "规则中的档位代码", "填写已发布规则中的工伤保险档位代码。"),
    ("social_base_mode", "社保基数模式", "条件必填", "启用社保且需要员工覆盖时", "公司默认/薪资全额/固定基数", "公司默认表示继承已发布规则；固定基数需填写对应金额。"),
    ("social_fixed_base", "社保固定基数", "条件必填", "社保基数模式为固定基数", "非负金额", "最终基数仍受已发布规则的上下限约束。"),
    ("fund_base_mode", "公积金基数模式", "条件必填", "启用公积金且需要员工覆盖时", "公司默认/薪资全额/固定基数", "公司默认表示继承已发布规则；固定基数需填写对应金额。"),
    ("fund_fixed_base", "公积金固定基数", "条件必填", "公积金基数模式为固定基数", "非负金额", "最终基数仍受已发布规则的上下限约束。"),
    ("tax_year", "个税年度", "条件必填", "填写累计个税数据时", "四位年份", "累计预扣数据所属自然年度。"),
    ("prior_taxable_income", "此前累计应税收入", "条件必填", "非首次计税或存在本年度历史累计数据时", "非负金额", "截至上一个计税月的累计应税收入；首次计税可留空。"),
    ("prior_special_deductions", "此前累计专项附加扣除", "条件必填", "非首次计税或存在本年度历史累计数据时", "非负金额", "截至上一个计税月的累计专项附加扣除；首次计税可留空。"),
    ("current_special_deductions", "本月专项附加扣除", "可选", "本月存在专项附加扣除时", "非负金额", "本计税月专项附加扣除金额；留空按 0 处理。"),
    ("prior_withheld_tax", "此前累计已预扣税额", "条件必填", "非首次计税或存在本年度历史累计数据时", "非负金额", "截至上一个计税月已预扣税额；首次计税可留空。"),
)

def _validate_field_guidance(headers, guidance):
    if tuple(item[0] for item in guidance) != tuple(headers):
        raise ValueError("员工薪酬字段说明与标准表头不一致")


_validate_field_guidance(EXPECTED_HEADERS["员工薪酬"], FIELD_GUIDANCE)


def build_employee_payroll_template() -> bytes:
    workbook = Workbook()
    guide = workbook.active
    guide.title = "填写说明"
    guide.append(["项目", "说明"])
    instructions = [
        ("填写流程", "先在 App 发布 HR 规则，再填写本模板，最后与当前计薪周期考勤文件一并上传计算。"),
        ("employee_id", "必填。与考勤文件工号完全一致；按文本填写以保留前导零。"),
        ("employment_type", "必填。填写已发布 HR 规则中的启用用工类型代码。"),
        ("薪资字段", "月薪、实习日薪/时薪及转正日期按用工类型填写。"),
        ("奖金与补贴", "填写当月金额；是否计税及是否进入缴费基数由已发布规则决定。"),
        ("累计个税", "非 1 月应填写本年度截至上月的累计收入、扣除及已扣税。"),
        ("数据安全", "模板、上传文件和计算结果仅保存在当前电脑本机；请按公司制度保管和备份。"),
        ("空白", "不适用字段留空；不要删除、改名或重复表头。"),
    ]
    for row in instructions:
        guide.append(row)
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 76
    guide.freeze_panes = "A2"

    guidance = workbook.create_sheet("字段说明")
    guidance.append(["标准字段", "中文名称", "填写要求", "适用条件", "格式或选项", "说明"])
    for row in FIELD_GUIDANCE:
        guidance.append(row)
    guidance.freeze_panes = "A2"
    guidance.auto_filter.ref = f"A1:F{guidance.max_row}"
    for column, width in zip("ABCDEF", (28, 24, 14, 34, 32, 62), strict=True):
        guidance.column_dimensions[column].width = width

    sheet = workbook.create_sheet("员工薪酬")
    headers = EXPECTED_HEADERS["员工薪酬"]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="17324D")
    input_fill = PatternFill("solid", fgColor="FFF3CD")
    for target in (guide, guidance, sheet):
        for cell in target[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    requirement_fills = {
        "必填": PatternFill("solid", fgColor="F8D7DA"),
        "条件必填": PatternFill("solid", fgColor="FFF3CD"),
        "可选": PatternFill("solid", fgColor="DDEBF7"),
    }
    for column, metadata in enumerate(FIELD_GUIDANCE, start=1):
        field, label, requirement, condition, value_format, description = metadata
        header = sheet.cell(1, column)
        header.fill = requirement_fills[requirement]
        header.font = Font(color="212529", bold=True)
        header.comment = Comment(
            f"{label}\n{requirement}\n适用条件：{condition}\n格式或选项：{value_format}\n{description}",
            "薪酬绩效工具",
        )
    for row in sheet.iter_rows(min_row=2, max_row=1001, max_col=len(headers)):
        for cell in row:
            cell.fill = input_fill
    _format_column(sheet, headers, "employee_id", "@")
    for name in ("hire_date", "termination_date", "conversion_date"):
        _format_column(sheet, headers, name, "yyyy-mm-dd")
    for name in headers:
        if name.endswith(("salary", "rate", "base", "income", "deductions", "tax")) or name in {"allowance", "bonus", "perfect_attendance_override"}:
            _format_column(sheet, headers, name, "#,##0.00")
    _validation(sheet, headers, "statutory_enabled", '"是,否"', allow_blank=False)
    _validation(sheet, headers, "social_base_mode", '"固定基数,薪资全额,公司默认"')
    _validation(sheet, headers, "fund_base_mode", '"固定基数,薪资全额,公司默认"')
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(15, min(28, len(header) + 4))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1001"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _format_column(sheet, headers, name, number_format):
    column = headers.index(name) + 1
    for row in range(2, 1002):
        sheet.cell(row, column).number_format = number_format


def _validation(sheet, headers, name, formula, *, allow_blank=True):
    column = get_column_letter(headers.index(name) + 1)
    validation = DataValidation(type="list", formula1=formula, allow_blank=allow_blank)
    validation.error = "请选择下拉列表中的值"
    validation.errorTitle = "填写值无效"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"{column}2:{column}1001")
