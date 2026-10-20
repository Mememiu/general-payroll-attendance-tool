"""Fixed-format HR payroll and policy workbook template."""

from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation


HR_WORKBOOK_SHEETS = (
    "填写说明",
    "基础与考勤",
    "考勤字段映射",
    "假勤与豁免",
    "五险一金",
    "员工薪酬",
    "个税累计",
)

BASE_FIELDS = (
    ("template_version", "模板版本", "1.2.0", "是", "请勿修改"),
    ("package_id", "规则包编号", None, "是", "公司规则唯一编号，例如 company-2026"),
    ("policy_version", "规则版本", None, "是", "每次修改规则后升级版本"),
    ("work_mode", "工时模式", None, "是", "固定工时或弹性工时"),
    ("flexible_rule_mode", "弹性工时规则", None, "否", "弹性工时模式下选择具体策略"),
    ("flexible_latest_start", "弹性最晚到岗时间", None, "否", "HH:MM；最晚到岗限制策略下必填"),
    ("core_start", "核心时段开始", None, "否", "HH:MM；核心时段覆盖策略下必填"),
    ("core_end", "核心时段结束", None, "否", "HH:MM；核心时段覆盖策略下必填"),
    ("daily_required_hours", "每日应出勤小时", None, "是", "由 HR 按公司制度填写，支持精确到分钟"),
    ("monthly_pay_days", "月平均计薪天数", None, "是", "由 HR 按公司制度填写"),
    ("standard_start", "标准上班时间", None, "是", "HH:MM"),
    ("standard_end", "标准下班时间", None, "是", "HH:MM"),
    ("lunch_start", "午休开始", None, "是", "HH:MM"),
    ("lunch_end", "午休结束", None, "是", "HH:MM"),
    ("earliest_valid_start", "最早有效打卡", None, "是", "HH:MM"),
    ("latest_normal_start", "最晚正常到岗", None, "是", "HH:MM"),
    ("late_grace_end", "迟到宽限截止", None, "是", "HH:MM"),
    ("grace_checkout", "宽限到岗下班边界", None, "否", "HH:MM；不用时留空"),
    ("checkout_boundary_mode", "下班边界模式", None, "否", "由 HR 按公司制度填写"),
    ("late_makeup_allowed", "迟到是否允许晚下班补足", None, "是", "是或否"),
    ("full_day_absence_after_hours", "迟到转整日旷工阈值", None, "是", "按小时填写"),
    ("attendance_deduction_factor", "考勤扣款系数", None, "是", "1 表示全额扣款"),
    ("sick_leave_deduction_factor", "病假扣款系数", None, "是", "0.3 表示扣 30%"),
    ("perfect_attendance_award", "公司全勤奖", None, "是", "金额，可由员工行覆盖"),
    ("overtime_cash_enabled", "加班是否发现金", None, "是", "是或否"),
    ("overtime_start_mode", "加班起算模式", None, "是", "选择加班开始时间的计算方式"),
    ("entry_day_full_attendance", "入职日是否按满勤", None, "是", "是或否"),
    ("termination_day_full_attendance", "离职日是否按满勤", None, "是", "是或否"),
    ("allowance_taxable", "补贴是否计税", None, "是", "是或否"),
    ("allowance_included_in_social_base", "补贴是否计入社保基数", None, "是", "是或否"),
    ("allowance_included_in_fund_base", "补贴是否计入公积金基数", None, "是", "是或否"),
    ("bonus_taxable", "奖金是否计税", None, "是", "是或否"),
    ("bonus_included_in_social_base", "奖金是否计入社保基数", None, "是", "是或否"),
    ("bonus_included_in_fund_base", "奖金是否计入公积金基数", None, "是", "是或否"),
)

LEAVE_ROWS = (
    ("sick_leave", "病假"),
    ("personal_leave", "事假"),
    ("compensatory_leave", "调休假"),
    ("paid_annual_leave", "带薪年假"),
    ("unpaid_annual_leave", "无薪年假"),
    ("bereavement_leave", "丧假"),
    ("external", "外出"),
    ("business_trip", "出差"),
    ("overtime", "加班"),
)

SOCIAL_ROWS = (
    ("pension", "养老保险"),
    ("medical", "医疗保险"),
    ("unemployment", "失业保险"),
    ("work_injury", "工伤保险"),
    ("maternity", "生育保险"),
    ("housing_fund", "住房公积金"),
)

EXPECTED_HEADERS = {
    "基础与考勤": ["field_code", "字段", "填写值", "必填", "说明"],
    "考勤字段映射": ["target_field", "标准字段", "必填", "上传文件列名", "说明"],
    "假勤与豁免": [
        "rule_code", "规则名称", "计薪比例", "计入出勤", "计入缺勤",
        "影响全勤", "免打卡", "需要证明", "备注",
    ],
    "五险一金": [
        "item_code", "险种", "档位代码", "启用", "地区", "生效日", "失效日",
        "基数模式", "基数下限", "基数上限", "固定基数", "个人费率", "单位费率",
        "政策文号", "来源网址", "HR复核日", "备注",
    ],
    "员工薪酬": [
        "employee_id", "statutory_enabled", "employee_name", "employment_type", "hire_date", "termination_date",
        "conversion_date", "monthly_salary", "intern_daily_rate", "intern_hourly_rate",
        "allowance", "bonus", "perfect_attendance_override", "pension_tier", "medical_tier",
        "work_injury_tier", "social_base_mode", "social_fixed_base", "fund_base_mode",
        "fund_fixed_base", "tax_year", "prior_taxable_income",
        "prior_special_deductions", "current_special_deductions",
        "prior_withheld_tax",
    ],
    "个税累计": [
        "employee_id", "tax_year", "prior_taxable_income", "prior_special_deductions",
        "prior_withheld_tax", "notes",
    ],
}


@dataclass(frozen=True)
class WorkbookValidationReport:
    errors: tuple[str, ...]
    template_version: str | None

    @property
    def ok(self) -> bool:
        return not self.errors


def build_hr_workbook_template() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _build_instructions(workbook.create_sheet("填写说明"))
    _build_base(workbook.create_sheet("基础与考勤"))
    _build_mapping(workbook.create_sheet("考勤字段映射"))
    _build_leave(workbook.create_sheet("假勤与豁免"))
    _build_social(workbook.create_sheet("五险一金"))
    _build_employee(workbook.create_sheet("员工薪酬"))
    _build_tax(workbook.create_sheet("个税累计"))
    workbook.properties.title = "薪酬与HR规则固定模板"
    workbook.properties.subject = "本地薪酬绩效工具输入模板"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def validate_hr_workbook(
    content: bytes, *, require_complete: bool = False
) -> WorkbookValidationReport:
    try:
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=False)
    except Exception as exc:
        return WorkbookValidationReport(("薪酬规则工作簿不是有效的 XLSX 文件",), None)
    try:
        errors: list[str] = []
        missing = set(HR_WORKBOOK_SHEETS) - set(workbook.sheetnames)
        if missing:
            errors.append(f"缺少工作表：{'、'.join(sorted(missing))}")
        extras = set(workbook.sheetnames) - set(HR_WORKBOOK_SHEETS)
        if extras:
            errors.append(f"包含不支持的工作表：{'、'.join(sorted(extras))}")
        for sheet_name, expected in EXPECTED_HEADERS.items():
            if sheet_name not in workbook.sheetnames:
                continue
            actual = [cell.value for cell in workbook[sheet_name][1]][:len(expected)]
            if actual != expected:
                errors.append(f"工作表 {sheet_name} 的表头已被修改")
        template_version = None
        if "基础与考勤" in workbook.sheetnames:
            template_version = workbook["基础与考勤"]["C2"].value
            if template_version != "1.2.0":
                errors.append("模板版本不受支持，请重新下载最新模板")
        if require_complete and not errors:
            errors.extend(_completeness_errors(workbook))
        return WorkbookValidationReport(tuple(errors), template_version)
    finally:
        workbook.close()


def _completeness_errors(workbook) -> list[str]:
    errors: list[str] = []
    base = workbook["基础与考勤"]
    missing_base = [
        base.cell(row, 2).value
        for row in range(2, base.max_row + 1)
        if base.cell(row, 4).value == "是" and _blank(base.cell(row, 3).value)
    ]
    if missing_base:
        errors.append(f"基础与考勤必填项未填写：{'、'.join(missing_base)}")
    base_values = {
        base.cell(row, 1).value: base.cell(row, 3).value
        for row in range(2, base.max_row + 1)
    }
    work_mode = base_values.get("work_mode")
    flexible_rule_mode = base_values.get("flexible_rule_mode")
    flexible_rule_modes = {"纯工时弹性", "最晚到岗限制", "核心时段覆盖"}
    if not _blank(work_mode) and work_mode not in {"固定工时", "弹性工时"}:
        errors.append("基础与考勤 work_mode 必须是固定工时或弹性工时")
    if work_mode == "固定工时" and not _blank(flexible_rule_mode):
        errors.append("基础与考勤 固定工时不得填写 flexible_rule_mode")
    if work_mode == "弹性工时":
        if flexible_rule_mode not in flexible_rule_modes:
            errors.append(
                "基础与考勤 flexible_rule_mode 必须是纯工时弹性、最晚到岗限制或核心时段覆盖"
            )
        elif flexible_rule_mode == "最晚到岗限制" and _blank(
            base_values.get("flexible_latest_start")
        ):
            errors.append("基础与考勤 最晚到岗限制必须填写 flexible_latest_start")
        elif flexible_rule_mode == "核心时段覆盖":
            for code in ("core_start", "core_end"):
                if _blank(base_values.get(code)):
                    errors.append(f"基础与考勤 核心时段覆盖必须填写 {code}")
    overtime_start_mode = base_values.get("overtime_start_mode")
    overtime_start_modes = {
        "审批开始时间", "正常下班后", "审批开始与正常下班取较晚",
    }
    if (
        not _blank(overtime_start_mode)
        and overtime_start_mode not in overtime_start_modes
    ):
        errors.append(
            "基础与考勤 overtime_start_mode 必须是审批开始时间、正常下班后或审批开始与正常下班取较晚"
        )
    for code in (
        "daily_required_hours", "monthly_pay_days", "full_day_absence_after_hours",
    ):
        if not _blank(base_values.get(code)) and not _positive_number(base_values[code]):
            errors.append(f"基础与考勤 {code} 必须是正数")
    for code in ("attendance_deduction_factor", "sick_leave_deduction_factor"):
        if not _blank(base_values.get(code)) and not _rate(base_values[code]):
            errors.append(f"基础与考勤 {code} 必须在 0 到 1 之间")
    for code in (
        "standard_start", "standard_end", "lunch_start", "lunch_end",
        "earliest_valid_start", "latest_normal_start", "late_grace_end",
        "grace_checkout", "flexible_latest_start", "core_start", "core_end",
    ):
        if not _blank(base_values.get(code)) and not _valid_time(base_values[code]):
            errors.append(f"基础与考勤 {code} 必须是 HH:MM 时间")

    mapping = workbook["考勤字段映射"]
    missing_mapping = [
        mapping.cell(row, 2).value
        for row in range(2, mapping.max_row + 1)
        if mapping.cell(row, 3).value == "是" and _blank(mapping.cell(row, 4).value)
    ]
    if missing_mapping:
        errors.append(f"考勤字段映射必填项未填写：{'、'.join(missing_mapping)}")

    leave = workbook["假勤与豁免"]
    incomplete_leave = [
        leave.cell(row, 2).value
        for row in range(2, leave.max_row + 1)
        if any(_blank(leave.cell(row, column).value) for column in range(3, 9))
    ]
    if incomplete_leave:
        errors.append(f"假勤与豁免规则未填写完整：{'、'.join(incomplete_leave)}")
    for row in range(2, leave.max_row + 1):
        ratio = leave.cell(row, 3).value
        if not _blank(ratio) and not _rate(ratio):
            errors.append(f"假勤与豁免 {leave.cell(row, 2).value} 的计薪比例必须在 0 到 1 之间")

    social = workbook["五险一金"]
    undecided_social = [
        social.cell(row, 2).value
        for row in range(2, social.max_row + 1)
        if social.cell(row, 4).value not in {"是", "否"}
    ]
    if undecided_social:
        errors.append(f"五险一金必须明确是否启用：{'、'.join(undecided_social)}")
    for row in range(2, social.max_row + 1):
        if social.cell(row, 4).value != "是":
            continue
        required_columns = (5, 6, 8, 9, 10, 12, 13, 14, 15, 16)
        missing = [
            social.cell(1, column).value
            for column in required_columns
            if _blank(social.cell(row, column).value)
        ]
        if social.cell(row, 8).value == "固定基数" and _blank(social.cell(row, 11).value):
            missing.append("固定基数")
        if missing:
            errors.append(
                f"五险一金 {social.cell(row, 2).value} 缺少：{'、'.join(missing)}"
            )
        for column, label in ((12, "个人费率"), (13, "单位费率")):
            value = social.cell(row, column).value
            if not _blank(value) and not _rate(value):
                errors.append(f"五险一金 {social.cell(row, 2).value} 的{label}必须在 0 到 1 之间")
        lower, upper = social.cell(row, 9).value, social.cell(row, 10).value
        if not _blank(lower) and not _non_negative_number(lower):
            errors.append(f"五险一金 {social.cell(row, 2).value} 的基数下限必须是非负数")
        if not _blank(upper) and not _non_negative_number(upper):
            errors.append(f"五险一金 {social.cell(row, 2).value} 的基数上限必须是非负数")
        if _non_negative_number(lower) and _non_negative_number(upper) and float(lower) > float(upper):
            errors.append(f"五险一金 {social.cell(row, 2).value} 的基数下限不能高于上限")

    employee = workbook["员工薪酬"]
    employee_rows = [
        row for row in range(2, employee.max_row + 1)
        if any(not _blank(employee.cell(row, column).value) for column in range(1, 21))
    ]
    if not employee_rows:
        errors.append("员工薪酬至少需要填写一名员工")
    employee_ids: set[str] = set()
    for row in employee_rows:
        employee_id = str(employee.cell(row, 1).value or "").strip()
        statutory_enabled = employee.cell(row, 2).value
        employment_type = employee.cell(row, 4).value
        if not employee_id:
            errors.append(f"员工薪酬第 {row} 行缺少 employee_id")
        elif employee_id in employee_ids:
            errors.append(f"员工薪酬工号重复：{employee_id}")
        employee_ids.add(employee_id)
        if statutory_enabled not in {"是", "否"}:
            errors.append(f"员工薪酬 {employee_id or row} 的是否缴纳五险一金必须填写是或否")
        if employment_type not in {"正职", "实习", "实习转正"}:
            errors.append(f"员工薪酬 {employee_id or row} 的用工类型无效")
        if employment_type in {"正职", "实习转正"} and _blank(employee.cell(row, 8).value):
            errors.append(f"员工薪酬 {employee_id or row} 缺少月薪")
        if employment_type in {"实习", "实习转正"} and (
            _blank(employee.cell(row, 9).value) or _blank(employee.cell(row, 10).value)
        ):
            errors.append(f"员工薪酬 {employee_id or row} 缺少实习日薪或时薪")
        if employment_type == "实习转正" and _blank(employee.cell(row, 7).value):
            errors.append(f"员工薪酬 {employee_id or row} 缺少转正日期")
        for column in range(8, 14):
            value = employee.cell(row, column).value
            if not _blank(value) and not _non_negative_number(value):
                errors.append(f"员工薪酬 {employee_id or row} 的金额必须是非负数")
        for mode_column, fixed_column, name in (
            (17, 18, "social_base_mode"),
            (19, 20, "fund_base_mode"),
        ):
            mode = employee.cell(row, mode_column).value
            fixed_base = employee.cell(row, fixed_column).value
            if not _blank(mode) and mode not in {"公司默认", "薪资全额", "固定基数"}:
                errors.append(
                    f"员工薪酬 {employee_id or row} 的{name}必须是公司默认、薪资全额或固定基数"
                )
            if mode == "固定基数" and _blank(fixed_base):
                errors.append(f"员工薪酬 {employee_id or row} 的{name}为固定基数时必须填写固定基数")
            if not _blank(fixed_base) and not _non_negative_number(fixed_base):
                errors.append(f"员工薪酬 {employee_id or row} 的{name}固定基数必须是非负数")

    tax = workbook["个税累计"]
    for row in range(2, tax.max_row + 1):
        if all(_blank(tax.cell(row, column).value) for column in range(1, 7)):
            continue
        if _blank(tax.cell(row, 1).value):
            errors.append(f"个税累计第 {row} 行缺少 employee_id")
        for column in (3, 4, 5):
            value = tax.cell(row, column).value
            if not _blank(value) and not _non_negative_number(value):
                errors.append(f"个税累计第 {row} 行金额必须是非负数")
    return errors


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _rate(value: object) -> bool:
    return _non_negative_number(value) and float(value) <= 1


def _positive_number(value: object) -> bool:
    return _non_negative_number(value) and float(value) > 0


def _non_negative_number(value: object) -> bool:
    return not isinstance(value, bool) and isinstance(value, (int, float)) and value >= 0


def _valid_time(value: object) -> bool:
    if isinstance(value, (time, datetime)):
        return True
    if not isinstance(value, str):
        return False
    try:
        datetime.strptime(value.strip(), "%H:%M")
        return True
    except ValueError:
        return False


def _build_instructions(sheet) -> None:
    sheet.append(["薪酬与 HR 规则固定模板", "版本 1.2.0"])
    instructions = (
        "请勿修改工作表名称、表头、field_code、target_field、rule_code 或 item_code。",
        "所有比例使用小数填写，例如 8% 填 0.08；金额填写数值，不输入货币符号。",
        "黄色单元格由 HR 填写；五险一金无地区默认值，必须依据当地有效政策手工填写并复核。",
        "员工工号是考勤与薪酬匹配主键，不允许为空或重复。",
        "上传时系统会校验模板版本、必填字段、日期、比例、档位和工号一致性。",
    )
    for text in instructions:
        sheet.append([text])
    sheet.column_dimensions["A"].width = 100
    sheet.column_dimensions["B"].width = 18
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["B1"].font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17324D")
    sheet.freeze_panes = "A2"


def _build_base(sheet) -> None:
    _append_headers(sheet, EXPECTED_HEADERS["基础与考勤"])
    for field in BASE_FIELDS:
        sheet.append(field)
    _input_cells(sheet, f"C2:C{sheet.max_row}")
    sheet["C2"].fill = PatternFill("solid", fgColor="E2E8F0")
    sheet["C2"].protection = Protection(locked=True)
    _list_validation(
        sheet,
        f"C{_row_for(sheet, 'work_mode')}",
        '"固定工时,弹性工时"',
    )
    _list_validation(
        sheet,
        f"C{_row_for(sheet, 'flexible_rule_mode')}",
        '"纯工时弹性,最晚到岗限制,核心时段覆盖"',
    )
    _list_validation(
        sheet,
        f"C{_row_for(sheet, 'overtime_start_mode')}",
        '"审批开始时间,正常下班后,审批开始与正常下班取较晚"',
    )
    for code in (
        "late_makeup_allowed", "overtime_cash_enabled", "entry_day_full_attendance",
        "termination_day_full_attendance", "allowance_taxable",
        "allowance_included_in_social_base", "allowance_included_in_fund_base",
        "bonus_taxable", "bonus_included_in_social_base", "bonus_included_in_fund_base",
    ):
        _list_validation(sheet, f"C{_row_for(sheet, code)}", '"是,否"')
    for code in (
        "standard_start", "standard_end", "lunch_start", "lunch_end",
        "earliest_valid_start", "latest_normal_start", "late_grace_end",
        "grace_checkout", "flexible_latest_start", "core_start", "core_end",
    ):
        sheet[f"C{_row_for(sheet, code)}"].number_format = "hh:mm"
    _finish_table(sheet, [34, 24, 20, 10, 52])


def _build_mapping(sheet) -> None:
    _append_headers(sheet, EXPECTED_HEADERS["考勤字段映射"])
    rows = (
        ("employee_id", "员工工号", "是", None, "用于与员工薪酬匹配"),
        ("work_date", "考勤日期", "是", None, "每日考勤所属日期"),
        ("attendance_text", "每日考勤结果", "是", None, "包含打卡与审批文本"),
        ("approval_reference", "关联审批单", "否", None, "仅作校验"),
    )
    for row in rows:
        sheet.append(row)
    _input_cells(sheet, f"D2:D{sheet.max_row}")
    _finish_table(sheet, [22, 22, 10, 28, 48])


def _build_leave(sheet) -> None:
    _append_headers(sheet, EXPECTED_HEADERS["假勤与豁免"])
    for code, name in LEAVE_ROWS:
        sheet.append([code, name, None, None, None, None, None, None, None])
    _input_cells(sheet, f"C2:I{sheet.max_row}")
    for column in "DEFGH":
        _list_validation(sheet, f"{column}2:{column}{sheet.max_row}", '"是,否"')
    for cell in sheet["C"][1:]:
        cell.number_format = "0.00%"
    _finish_table(sheet, [24, 18, 14, 14, 14, 14, 12, 14, 36])


def _build_social(sheet) -> None:
    _append_headers(sheet, EXPECTED_HEADERS["五险一金"])
    for code, name in SOCIAL_ROWS:
        sheet.append([code, name] + [None] * 15)
    _input_cells(sheet, f"C2:Q{sheet.max_row}")
    _list_validation(sheet, f"D2:D{sheet.max_row}", '"是,否"')
    _list_validation(sheet, f"H2:H{sheet.max_row}", '"固定基数,薪资全额,员工单独基数"')
    for column in "FGP":
        for cell in sheet[column][1:]:
            cell.number_format = "yyyy-mm-dd"
    for column in "LM":
        for cell in sheet[column][1:]:
            cell.number_format = "0.0000%"
    for column in "IJK":
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0.00"
    _finish_table(sheet, [20, 18, 16, 10, 18, 14, 14, 20, 14, 14, 14, 14, 14, 20, 32, 14, 34])


def _build_employee(sheet) -> None:
    _append_headers(sheet, EXPECTED_HEADERS["员工薪酬"])
    _input_cells(sheet, "A2:S1001")
    _list_validation(sheet, "B2:B1001", '"是,否"')
    _list_validation(sheet, "D2:D1001", '"正职,实习,实习转正"')
    _list_validation(sheet, "Q2:Q1001", '"固定基数,薪资全额,公司默认"')
    _list_validation(sheet, "S2:S1001", '"固定基数,薪资全额,公司默认"')
    for column in "EFG":
        for cell in sheet[column][1:]:
            cell.number_format = "yyyy-mm-dd"
    for column in "HIJKLMRT":
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0.00"
    _finish_table(sheet, [18, 16, 16, 14, 14, 14] + [16] * 13)


def _build_tax(sheet) -> None:
    _append_headers(sheet, EXPECTED_HEADERS["个税累计"])
    _input_cells(sheet, "A2:F1001")
    for column in "CDE":
        for cell in sheet[column][1:]:
            cell.number_format = "#,##0.00"
    _finish_table(sheet, [18, 12, 20, 20, 18, 32])


def _append_headers(sheet, headers: list[str]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.protection = Protection(locked=True)


def _input_cells(sheet, cell_range: str) -> None:
    fill = PatternFill("solid", fgColor="FFF3CD")
    for row in sheet[cell_range]:
        cells = row if isinstance(row, tuple) else (row,)
        for cell in cells:
            cell.fill = fill
            cell.protection = Protection(locked=False)


def _list_validation(sheet, cell_range: str, formula: str) -> None:
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "请选择下拉列表中的有效值"
    validation.errorTitle = "填写值无效"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def _row_for(sheet, code: str) -> int:
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == code:
            return row
    raise ValueError(f"unknown template field: {code}")


def _finish_table(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 40
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.comment = Comment("固定表头，请勿修改", "Payroll Tool")
