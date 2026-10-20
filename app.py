"""Offline HR workspace for current-task payroll uploads."""

from datetime import date
from datetime import datetime, timezone
from dataclasses import asdict, replace
from decimal import Decimal
from io import BytesIO
from hashlib import sha256
import json
from pathlib import Path
import secrets
import shutil
import tempfile
from uuid import uuid4

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import load_workbook
from werkzeug.datastructures import FileStorage
from werkzeug.utils import secure_filename

from payroll_tool.templates.hr_workbook import build_hr_workbook_template, validate_hr_workbook
from payroll_tool.templates.employee_workbook import build_employee_payroll_template
from payroll_tool.importers.excel import import_attendance_workbook
from payroll_tool.importers.hr_workbook import (
    CumulativeTaxInput,
    attendance_policy_from_hr,
    income_item_policies_from_hr,
    parse_hr_workbook,
)
from payroll_tool.importers.uploaded_inputs import CalculationInputSession
from payroll_tool.importers.employee_workbook import parse_employee_workbook
from payroll_tool.config.policy_serialization import calculation_policy_from_snapshot
from payroll_tool.attendance.batch import calculate_attendance_rows
from payroll_tool.payroll.engine import calculate_payroll
from payroll_tool.statutory.social_insurance import applicable_statutory_tiers
from payroll_tool.audit import build_audit_manifest
from payroll_tool.exporters.excel import build_payroll_export
from payroll_tool.rule_packs import RulePackStore
from payroll_tool.config.policy_repository import PolicyRepository
from payroll_tool.web.policy_api import policy_api
from payroll_tool.web.backup_api import backup_api


REQUIRED_UPLOADS = {
    "attendance": ("考勤文件", {".xlsx"}),
    "payroll_rules": ("薪酬规则工作簿", {".xlsx"}),
}


class TaskStore:
    def __init__(self, root: str | Path | None = None) -> None:
        self._temporary = tempfile.TemporaryDirectory(prefix="general-payroll-") if root is None else None
        self._root_path = Path(self._temporary.name) if self._temporary else Path(root)
        self._root_path.mkdir(parents=True, exist_ok=True, mode=0o700)
        self._root_path.chmod(0o700)
        self.tasks: dict[str, dict[str, object]] = {}

    def create(self, uploads: dict[str, object]) -> tuple[str, str]:
        task_id = uuid4().hex
        token = secrets.token_urlsafe(24)
        task_directory = self._root_path / task_id
        task_directory.mkdir(mode=0o700)
        paths: dict[str, Path] = {}
        for role, upload in uploads.items():
            suffix = Path(secure_filename(upload.filename)).suffix.lower() or ".upload"
            filename = f"{role}{suffix}"
            path = task_directory / filename
            upload.save(path)
            paths[role] = path
        self.tasks[task_id] = {"token": token, "paths": paths, "status": "uploaded"}
        self._write_metadata(task_id, {"token_sha256": sha256(token.encode()).hexdigest()})
        return task_id, token

    def save_result(self, task_id: str, content: bytes) -> Path:
        result = self._root_path / task_id / "result.xlsx"
        result.write_bytes(content)
        result.chmod(0o600)
        return result

    def authorized_result(self, task_id: str, token: str) -> Path | None:
        try:
            metadata = json.loads((self._root_path / task_id / "metadata.json").read_text())
        except (OSError, ValueError, TypeError):
            return None
        supplied = sha256(token.encode()).hexdigest()
        result = self._root_path / task_id / "result.xlsx"
        if not secrets.compare_digest(metadata.get("token_sha256", ""), supplied):
            return None
        return result if result.is_file() else None

    def _write_metadata(self, task_id: str, value: dict[str, object]) -> None:
        path = self._root_path / task_id / "metadata.json"
        path.write_text(json.dumps(value, sort_keys=True), encoding="utf-8")
        path.chmod(0o600)

    def close(self) -> None:
        self.tasks.clear()
        if self._temporary is not None:
            self._temporary.cleanup()

    @property
    def storage_root(self) -> Path:
        return self._root_path

    def discard(self, task_id: str) -> None:
        task = self.tasks.pop(task_id, None)
        if task is not None:
            task_directory = next(iter(task["paths"].values())).parent
            shutil.rmtree(task_directory, ignore_errors=True)


def create_app(config: dict | None = None) -> Flask:
    app = Flask(__name__)
    app.config.update(
        MAX_CONTENT_LENGTH=30 * 1024 * 1024,
        RULE_PACK_ROOT=Path(app.instance_path) / "rule-packs",
        POLICY_DATABASE=Path(app.instance_path) / "policies" / "policies.sqlite3",
        TASK_ROOT=Path(app.instance_path) / "tasks",
    )
    if config:
        app.config.update(config)
    task_store = TaskStore(None if app.config.get("TESTING") and not (config and "TASK_ROOT" in config) else app.config["TASK_ROOT"])
    if app.config.get("TESTING") and not (config and "POLICY_DATABASE" in config):
        app.config["POLICY_DATABASE"] = task_store.storage_root / "policies.sqlite3"
    app.extensions["task_store"] = task_store
    app.extensions["policy_repository"] = PolicyRepository(app.config["POLICY_DATABASE"])
    app.register_blueprint(policy_api)
    app.register_blueprint(backup_api)

    @app.get("/")
    def workspace():
        return render_template("workspace.html")

    @app.get("/api/calculation-records")
    def calculation_records():
        repository = app.extensions["policy_repository"]
        return jsonify(
            items=[record.as_dict() for record in repository.list_calculation_records()]
        )

    @app.get("/templates/payroll-rules.xlsx")
    def payroll_rules_template():
        return send_file(
            BytesIO(build_hr_workbook_template()),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="薪酬与HR规则完整模板.xlsx",
        )

    @app.get("/templates/employee-payroll.xlsx")
    def employee_payroll_template():
        return send_file(
            BytesIO(build_employee_payroll_template()),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="员工薪酬模板.xlsx",
        )

    @app.get("/rule-packs")
    def list_rule_packs():
        store = _rule_pack_store(app)
        return jsonify({"items": [asdict(pack) for pack in store.list()]})

    @app.post("/rule-packs")
    def save_rule_pack():
        errors: list[str] = []
        name = request.form.get("name", "").strip()
        version = request.form.get("version", "").strip()
        effective_from = _parse_date(request.form.get("effective_from", ""), "生效日", errors)
        upload = request.files.get("workbook")
        if not name:
            errors.append("请输入规则包名称")
        if not version:
            errors.append("请输入规则版本")
        if upload is None or not upload.filename:
            errors.append("请上传规则工作簿")
        elif Path(upload.filename).suffix.lower() != ".xlsx":
            errors.append("规则工作簿仅支持 XLSX")
        else:
            content = upload.read()
            report = validate_hr_workbook(content, require_complete=True)
            errors.extend(report.errors)
        if errors:
            return jsonify({"ok": False, "errors": errors}), 400
        pack = _rule_pack_store(app).save(
            name=name,
            version=version,
            effective_from=effective_from,
            workbook_bytes=content,
        )
        return jsonify({"ok": True, **asdict(pack)}), 201

    @app.get("/rule-packs/<pack_id>/download")
    def download_rule_pack(pack_id: str):
        store = _rule_pack_store(app)
        try:
            pack = store.get(pack_id)
            content = store.read_workbook(pack_id)
        except ValueError:
            return jsonify({"ok": False, "errors": ["规则包不存在或已损坏"]}), 404
        return send_file(
            BytesIO(content),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"{secure_filename(pack.name)}-{secure_filename(pack.version)}.xlsx",
        )

    @app.post("/rule-packs/<pack_id>/clone")
    def clone_rule_pack(pack_id: str):
        errors: list[str] = []
        version = request.form.get("version", "").strip()
        effective_from = _parse_date(request.form.get("effective_from", ""), "生效日", errors)
        upload = request.files.get("workbook")
        store = _rule_pack_store(app)
        try:
            store.get(pack_id)
        except ValueError:
            errors.append("源规则包不存在")
        if not version:
            errors.append("请输入新规则版本")
        if upload is None or not upload.filename:
            errors.append("请上传更新后的规则工作簿")
        elif Path(upload.filename).suffix.lower() != ".xlsx":
            errors.append("规则工作簿仅支持 XLSX")
        else:
            content = upload.read()
            report = validate_hr_workbook(content, require_complete=True)
            errors.extend(report.errors)
        if errors:
            return jsonify({"ok": False, "errors": errors}), 400
        pack = store.clone(
            pack_id,
            version=version,
            effective_from=effective_from,
            workbook_bytes=content,
        )
        return jsonify({"ok": True, **asdict(pack)}), 201

    @app.post("/tasks")
    def create_task():
        uploads: dict[str, object] = {}
        errors: list[str] = []
        rule_pack_id = request.form.get("rule_pack_id", "").strip()
        policy_version_id = request.form.get("policy_version_id", "").strip()
        policy_version = None
        payroll_month = request.form.get("payroll_month", "").strip()
        period_start = _parse_date(request.form.get("period_start", ""), "起薪日", errors)
        period_end = _parse_date(request.form.get("period_end", ""), "止薪日", errors)
        if len(payroll_month) != 7 or payroll_month[4:5] != "-":
            errors.append("请选择有效的计薪月份")
        if period_start is not None and period_end is not None and period_start > period_end:
            errors.append("起薪日不能晚于止薪日")
        if policy_version_id:
            try:
                policy_version = app.extensions["policy_repository"].get(policy_version_id)
                if policy_version.status.value != "published":
                    errors.append("请选择已发布的规则版本")
            except ValueError:
                errors.append("所选规则版本不存在")
        if rule_pack_id:
            try:
                rule_bytes = _rule_pack_store(app).read_workbook(rule_pack_id)
                report = validate_hr_workbook(rule_bytes, require_complete=True)
                errors.extend(report.errors)
                uploads["payroll_rules"] = FileStorage(
                    stream=BytesIO(rule_bytes), filename=f"rule-pack-{rule_pack_id}.xlsx"
                )
            except ValueError:
                errors.append("所选规则包不存在或已损坏")
        required_uploads = (
            {"attendance": REQUIRED_UPLOADS["attendance"]}
            if policy_version is not None else REQUIRED_UPLOADS
        )
        for role, (label, extensions) in required_uploads.items():
            if role == "payroll_rules" and rule_pack_id:
                continue
            upload = request.files.get(role)
            if upload is None or not upload.filename:
                errors.append(f"请上传{label}")
                continue
            if Path(upload.filename).suffix.lower() not in extensions:
                errors.append(f"{label}文件格式不支持")
                continue
            if role == "payroll_rules":
                content = upload.read()
                report = validate_hr_workbook(content, require_complete=True)
                errors.extend(report.errors)
                upload.seek(0)
                if not report.ok:
                    continue
            uploads[role] = upload
        employee_payroll = request.files.get("employee_payroll")
        if policy_version is not None and (
            employee_payroll is None or not employee_payroll.filename
        ):
            errors.append("请上传员工薪酬文件")
        if employee_payroll is not None and employee_payroll.filename:
            if Path(employee_payroll.filename).suffix.lower() != ".xlsx":
                errors.append("员工薪酬文件格式不支持")
            else:
                uploads["employee_payroll"] = employee_payroll
        if errors:
            return jsonify({"ok": False, "errors": errors}), 400
        task_id, token = app.extensions["task_store"].create(uploads)
        task = app.extensions["task_store"].tasks[task_id]
        record_created = False
        if policy_version is not None:
            payroll_input = (
                task["paths"]["employee_payroll"]
                if "employee_payroll" in task["paths"]
                else task["paths"]["payroll_rules"]
            )
            app.extensions["policy_repository"].create_calculation_record(
                task_id=task_id,
                policy_version_id=policy_version.id,
                policy_digest=policy_version.snapshot_sha256,
                payroll_month=payroll_month,
                period_start=period_start.isoformat(),
                period_end=period_end.isoformat(),
                attendance_digest=_file_digest(task["paths"]["attendance"]),
                employee_payroll_digest=_file_digest(payroll_input),
            )
            record_created = True
        try:
            if policy_version is not None:
                calculation_policy = calculation_policy_from_snapshot(
                    policy_version.snapshot
                )
                employees_list = []
                employment_policies = {}
                active_exception_count = 0

                def apply_period_exceptions(employee):
                    nonlocal active_exception_count
                    period_exceptions = (
                        calculation_policy.employee_exceptions_for_period(
                            employee.employee_id, period_start, period_end
                        )
                    )
                    _validate_period_exception_coverage(
                        period_exceptions, period_start, period_end
                    )
                    exceptions = tuple(
                        item for item in period_exceptions
                        if item.override_type == "employment_type"
                        or (
                            item.effective_from <= period_start
                            and (
                                item.effective_to is None
                                or period_start <= item.effective_to
                            )
                        )
                    )
                    active_exception_count += len(exceptions)
                    return _apply_employee_exceptions(employee, exceptions)

                def resolve_statutory_tiers(employee):
                    employment = calculation_policy.employment_types[
                        employee.employment_type
                    ]
                    if not employment.social_insurance_applicable:
                        return None
                    rules = tuple(
                        rule for rule in calculation_policy.statutory_rules
                        if rule.item_code != "housing_fund"
                        and (
                            not rule.employment_types
                            or employment.code in rule.employment_types
                        )
                    )
                    return applicable_statutory_tiers(
                        employee,
                        rules,
                        period_date=date.fromisoformat(f"{payroll_month}-01"),
                        effective_date=period_start,
                    )

                for employee in parse_employee_workbook(
                    task["paths"]["employee_payroll"].read_bytes(),
                    calculation_policy.employment_types,
                    resolve_statutory_tiers,
                    employee_transform=apply_period_exceptions,
                ):
                    employment = calculation_policy.employment_types[employee.employment_type]
                    social_enabled = (
                        employee.social_insurance_enabled
                        if employee.social_insurance_enabled is not None
                        else employee.statutory_enabled
                    ) and employment.social_insurance_applicable
                    housing_enabled = (
                        employee.housing_fund_enabled
                        if employee.housing_fund_enabled is not None
                        else employee.statutory_enabled
                    ) and employment.housing_fund_applicable
                    employee = replace(
                        employee,
                        statutory_enabled=social_enabled or housing_enabled,
                        social_insurance_enabled=social_enabled,
                        housing_fund_enabled=housing_enabled,
                        employment_type={
                        "full_time_monthly": "full_time",
                        "intern_daily_hourly": "intern",
                        "conversion_split": "converted",
                        }[employment.payroll_scheme],
                    )
                    employees_list.append(employee)
                    employment_policies[employee.employee_id] = employment
                employees = tuple(employees_list)
                attendance_mapping = {
                    "employee_id": "工号", "work_date": "日期",
                    "attendance_text": "每日考勤结果",
                }
                if "工作日类型" in _workbook_headers(task["paths"]["attendance"]):
                    attendance_mapping["workday_type"] = "工作日类型"
                attendance_policy = calculation_policy.attendance
                overtime_policy = calculation_policy.overtime
                leave_rules = calculation_policy.leave_rules
                income_policies = calculation_policy.income_items
                statutory_rules = calculation_policy.statutory_rules
                commercial_insurance = calculation_policy.commercial_insurance
                company_perfect_attendance_award = calculation_policy.perfect_attendance_award
                configured_daily_hours = Decimal(calculation_policy.daily_minutes) / Decimal("60")
                configured_monthly_days = calculation_policy.monthly_pay_days
                task["config_bytes"] = json.dumps(policy_version.snapshot, ensure_ascii=False, sort_keys=True).encode()
                task["rule_package_id"] = policy_version.policy_id
                task["rule_version"] = policy_version.version
            else:
                employment_policies = {}
                active_exception_count = 0
                parsed_rules = parse_hr_workbook(task["paths"]["payroll_rules"].read_bytes())
                employees = parsed_rules.employees
                attendance_mapping = parsed_rules.attendance_mapping
                attendance_policy = attendance_policy_from_hr(parsed_rules)
                leave_rules = parsed_rules.leave_rules
                income_policies = income_item_policies_from_hr(parsed_rules)
                statutory_rules = parsed_rules.statutory_rules
                commercial_insurance = ()
                company_perfect_attendance_award = Decimal(
                    str(parsed_rules.base_values["perfect_attendance_award"])
                )
                configured_daily_hours = Decimal(str(parsed_rules.base_values["daily_required_hours"]))
                configured_monthly_days = Decimal(str(parsed_rules.base_values["monthly_pay_days"]))
                overtime_policy = None
                task["config_bytes"] = task["paths"]["payroll_rules"].read_bytes()
                task["rule_package_id"] = str(parsed_rules.base_values["package_id"])
                task["rule_version"] = str(parsed_rules.base_values["policy_version"])
            session = CalculationInputSession()
            attendance_upload = session.register_attendance(task["paths"]["attendance"])
            attendance = import_attendance_workbook(
                attendance_upload, attendance_mapping
            )
            attendance_rows = _normalize_workday_types(attendance.rows)
            summaries = calculate_attendance_rows(
                attendance_rows,
                {employee.employee_id for employee in employees},
                period_start,
                period_end,
                attendance_policy,
                leave_rules,
                employee_periods={
                    employee.employee_id: (employee.hire_date, employee.termination_date)
                    for employee in employees
                },
                comp_time_ratio=(
                    overtime_policy.comp_time_ratio
                    if overtime_policy is not None
                    and overtime_policy.comp_time_enabled
                    else None
                ),
            )
            if overtime_policy is not None and overtime_policy.cash_enabled:
                summaries = _review_unsettleable_cash_overtime(summaries)
            statutory_period = date.fromisoformat(f"{payroll_month}-01")
            payroll_results = {}
            for employee in employees:
                summary = summaries.get(employee.employee_id)
                employment_policy = employment_policies.get(employee.employee_id)
                employee_statutory_rules = statutory_rules
                if employment_policy is not None:
                    employee_statutory_rules = tuple(
                        rule for rule in statutory_rules
                        if (
                            employment_policy.housing_fund_applicable
                            if rule.item_code == "housing_fund"
                            else employment_policy.social_insurance_applicable
                        )
                        and (
                            not rule.employment_types
                            or employment_policy.code in rule.employment_types
                        )
                    )
                if policy_version is None:
                    cumulative_tax = parsed_rules.cumulative_tax.get(
                        employee.employee_id,
                        CumulativeTaxInput(
                            tax_year=statutory_period.year,
                            prior_taxable_income=Decimal("0"),
                            prior_special_deductions=Decimal("0"),
                            prior_withheld_tax=Decimal("0"),
                        ),
                    )
                    tax_monthly_exemption = Decimal("5000")
                    tax_brackets = None
                elif (
                    calculation_policy.tax.enabled
                    and calculation_policy.tax.effective_from <= period_start
                    and employment_policy.tax_applicable
                ):
                    cumulative_tax = CumulativeTaxInput(
                        tax_year=employee.tax_year or statutory_period.year,
                        prior_taxable_income=employee.prior_taxable_income,
                        prior_special_deductions=employee.prior_special_deductions,
                        prior_withheld_tax=employee.prior_withheld_tax,
                    )
                    tax_monthly_exemption = calculation_policy.tax.monthly_exemption
                    tax_brackets = tuple(
                        (item.threshold, item.rate, item.quick_deduction)
                        for item in calculation_policy.tax.brackets
                    )
                    negative_tax_mode = calculation_policy.tax.negative_tax_mode
                else:
                    cumulative_tax = None
                    tax_monthly_exemption = Decimal("5000")
                    tax_brackets = None
                    negative_tax_mode = "zero"
                if policy_version is None:
                    negative_tax_mode = "zero"
                payroll_results[employee.employee_id] = calculate_payroll(
                    employee,
                    summary.attendance_days if summary else [],
                    income_items=income_policies,
                    statutory_rules=employee_statutory_rules,
                    period_date=statutory_period,
                    cumulative_tax=cumulative_tax,
                    perfect_attendance_award=company_perfect_attendance_award,
                    perfect_attendance_eligible=bool(
                        summary
                        and summary.processed_days > 0
                        and summary.absence_hours == 0
                        and summary.late_hours == 0
                        and not summary.has_leave
                    ),
                    daily_required_hours=configured_daily_hours,
                    monthly_pay_days=configured_monthly_days,
                    tax_monthly_exemption=tax_monthly_exemption,
                    tax_brackets=tax_brackets,
                    negative_tax_mode=negative_tax_mode,
                    statutory_effective_date=period_start,
                    current_special_deductions=employee.current_special_deductions,
                    commercial_insurance=commercial_insurance,
                    commercial_employment_type=(
                        employment_policy.code if employment_policy is not None else None
                    ),
                    overtime_policy=overtime_policy,
                )
        except ValueError as exc:
            if record_created:
                app.extensions["policy_repository"].fail_calculation_record(
                    task_id, str(exc)
                )
            app.extensions["task_store"].discard(task_id)
            return jsonify({"ok": False, "errors": [str(exc)]}), 400
        except Exception:
            if record_created:
                app.extensions["policy_repository"].fail_calculation_record(
                    task_id, "内部计算错误"
                )
            app.extensions["task_store"].discard(task_id)
            raise
        task["status"] = "validated"
        task["employee_count"] = len(employees)
        task["attendance_row_count"] = len(attendance.rows)
        task["processed_day_count"] = sum(item.processed_days for item in summaries.values())
        task["review_item_count"] = sum(len(item.review_reasons) for item in summaries.values())
        task["attendance_summaries"] = summaries
        task["employee_names"] = {
            employee.employee_id: employee.employee_name for employee in employees
        }
        task["payroll_results"] = payroll_results
        try:
            app.extensions["task_store"].save_result(task_id, _build_task_export(task))
            if record_created:
                review_count = task["review_item_count"]
                app.extensions["policy_repository"].complete_calculation_record(
                    task_id,
                    employee_count=len(employees),
                    exception_count=review_count + active_exception_count,
                    result_available=True,
                    needs_review=review_count > 0,
                )
        except Exception:
            if record_created:
                app.extensions["policy_repository"].fail_calculation_record(
                    task_id, "结果文件写入失败"
                )
            app.extensions["task_store"].discard(task_id)
            raise
        return jsonify(
            {
                "ok": True,
                "task_id": task_id,
                "download_token": token,
                "employee_count": task["employee_count"],
                "attendance_row_count": task["attendance_row_count"],
                "processed_day_count": task["processed_day_count"],
                "review_item_count": task["review_item_count"],
                "payroll_result_count": len(payroll_results),
            }
        ), 201

    @app.get("/tasks/<task_id>/download")
    def download_task(task_id: str):
        token = request.args.get("token", "")
        result = app.extensions["task_store"].authorized_result(task_id, token)
        if result is None:
            return jsonify({"ok": False, "errors": ["下载令牌无效"]}), 403
        return send_file(
            BytesIO(result.read_bytes()),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"薪酬计算结果-{task_id[:8]}.xlsx",
        )

    return app


def create_desktop_app(data_root=None) -> Flask:
    from payroll_tool.desktop.migrations import ensure_data_version
    from payroll_tool.desktop.storage import production_config, user_data_root

    root = Path(data_root) if data_root is not None else user_data_root()
    ensure_data_version(root)
    return create_app(production_config(root))


def _parse_date(value: str, label: str, errors: list[str]) -> date | None:
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        errors.append(f"请选择有效的{label}")
        return None


def _rule_pack_store(app: Flask) -> RulePackStore:
    return RulePackStore(app.config["RULE_PACK_ROOT"])


def _file_digest(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _build_task_export(task: dict[str, object]) -> bytes:
    rows = []
    for employee_id, payroll_result in task["payroll_results"].items():
        summary = task.get("attendance_summaries", {}).get(employee_id)
        rows.append({
            "employee_id": employee_id,
            "employee_name": task.get("employee_names", {}).get(employee_id, ""),
            "attendance_days": summary.processed_days if summary else 0,
            "effective_hours": summary.effective_hours if summary else Decimal("0"),
            "absence_hours": summary.absence_hours if summary else Decimal("0"),
            "late_hours": summary.late_hours if summary else Decimal("0"),
            "raw_overtime_hours": (
                summary.raw_overtime_minutes / Decimal("60")
                if summary else Decimal("0")
            ),
            "overtime_hours": (
                summary.overtime_minutes / Decimal("60")
                if summary else Decimal("0")
            ),
            "comp_time_accrual_minutes": (
                summary.comp_time_accrual_minutes if summary else Decimal("0")
            ),
            "leave_hours": json.dumps({kind: str(hours) for kind, hours in (summary.leave_hours if summary else {}).items()}, ensure_ascii=False, sort_keys=True),
            "review_reasons": "；".join(summary.review_reasons if summary else ()),
            **asdict(payroll_result),
        })
    result_bytes = json.dumps(rows, default=str, ensure_ascii=True, sort_keys=True).encode()
    audit = build_audit_manifest(
        input_bytes=task["paths"]["attendance"].read_bytes(),
        config_bytes=task["config_bytes"], result_bytes=result_bytes,
        timestamp=datetime.now(timezone.utc),
        rule_package_id=task["rule_package_id"], rule_version=task["rule_version"],
    )
    return build_payroll_export(rows, audit)


def _workbook_headers(path: Path) -> set[str]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(
            "uploaded attendance workbook is not a valid .xlsx file"
        ) from exc
    try:
        first_row = next(workbook.active.iter_rows(values_only=True), ())
        return {
            str(value).strip() for value in first_row
            if value is not None and str(value).strip()
        }
    finally:
        workbook.close()


def _normalize_workday_types(
    rows: tuple[dict[str, object], ...],
) -> tuple[dict[str, object], ...]:
    aliases = {
        "工作日": "workday",
        "休息日": "rest_day",
        "法定节假日": "holiday",
    }
    normalized = []
    for row in rows:
        item = dict(row)
        raw_value = item.get("workday_type")
        item["workday_type"] = aliases.get(str(raw_value).strip()) if raw_value else None
        normalized.append(item)
    return tuple(normalized)


def _review_unsettleable_cash_overtime(summaries):
    reviewed = {}
    for employee_id, summary in summaries.items():
        reasons = list(summary.review_reasons)
        attendance_days = []
        for day in summary.attendance_days:
            if day.overtime_minutes > 0 and day.workday_type is None:
                reasons.append(
                    f"{day.work_date.isoformat()} 认可现金加班缺少有效工作日类型，"
                    "加班费暂按 0 并需人工复核"
                )
                day = replace(day, overtime_minutes=Decimal("0"))
            attendance_days.append(day)
        reviewed[employee_id] = replace(
            summary,
            review_reasons=tuple(reasons),
            attendance_days=tuple(attendance_days),
        )
    return reviewed


def _apply_employee_exceptions(employee, exceptions):
    updates = {}
    tier_fields = {
        "pension": "pension_tier", "medical": "medical_tier",
        "work_injury": "work_injury_tier",
    }
    base_modes = {"fixed": "固定基数", "salary_full": "薪资全额", "company_default": None}
    for item in exceptions:
        kind = item.override_type
        if kind == "perfect_attendance_amount":
            updates["perfect_attendance_override"] = item.value
        elif kind == "employment_type":
            updates["employment_type"] = item.value
        elif kind == "statutory_enabled":
            updates["statutory_enabled"] = item.value
        elif kind == "statutory_tier":
            field = tier_fields.get(item.target_code)
            if field:
                updates[field] = item.value
            overrides = dict(updates.get(
                "statutory_tier_overrides", employee.statutory_tier_overrides
            ))
            overrides[item.target_code] = item.value
            updates["statutory_tier_overrides"] = tuple(sorted(overrides.items()))
        elif kind == "statutory_base_mode":
            updates["social_base_mode"] = base_modes[item.value]
        elif kind == "statutory_fixed_base":
            updates["social_fixed_base"] = item.value
        elif kind == "housing_fund_enabled":
            updates["housing_fund_enabled"] = item.value
        elif kind == "housing_fund_tier":
            overrides = dict(updates.get(
                "statutory_tier_overrides", employee.statutory_tier_overrides
            ))
            overrides["housing_fund"] = item.value
            updates["statutory_tier_overrides"] = tuple(sorted(overrides.items()))
        elif kind == "housing_fund_base_mode":
            updates["fund_base_mode"] = base_modes[item.value]
        elif kind == "housing_fund_fixed_base":
            updates["fund_fixed_base"] = item.value
        elif kind in {"commercial_enabled", "commercial_amount", "commercial_base"}:
            field = {
                "commercial_enabled": "commercial_enabled_overrides",
                "commercial_amount": "commercial_amount_overrides",
                "commercial_base": "commercial_base_overrides",
            }[kind]
            overrides = dict(updates.get(field, getattr(employee, field)))
            overrides[item.target_code] = item.value
            updates[field] = tuple(sorted(overrides.items()))
    return replace(employee, **updates)


def _validate_period_exception_coverage(exceptions, period_start, period_end):
    for item in exceptions:
        if item.override_type != "employment_type":
            continue
        covers_period = (
            item.effective_from <= period_start
            and (item.effective_to is None or item.effective_to >= period_end)
        )
        if not covers_period:
            raise ValueError(
                "用工类型例外必须覆盖完整计薪周期；月中变更请使用实习转正方案"
            )


if __name__ == "__main__":
    create_app().run(host="127.0.0.1", port=5050, debug=False)
